from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "Atelier_Pratique_Pipeline_STM32Cube_improved.pptx"
OUTPUT = ROOT / "STM32Cube_Preprocessing_Solution_40min_English.pptx"
TEMPLATE_CANDIDATES = [
    ROOT / "Atelier_Pratique_Pipeline_STM32Cube_improved.pptx",
    ROOT / "Atelier_Pratique_Pipeline_STM32Cube_working_copy.pptx",
    ROOT / "Atelier_Pratique_Pipeline_STM32Cube.pptx",
    ROOT / "Atelier_Pratique_Pipeline_STM32Cube_updated.pptx",
]

ST_BLUE = RGBColor(26, 35, 126)
ST_LIGHT_BLUE = RGBColor(218, 229, 255)
ST_ORANGE = RGBColor(255, 152, 0)
ST_GREEN = RGBColor(76, 175, 80)
ST_RED = RGBColor(210, 60, 60)
ST_GRAY = RGBColor(92, 96, 105)
LIGHT_GRAY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(22, 24, 29)
PURPLE = RGBColor(92, 82, 180)
CYAN = RGBColor(0, 135, 170)

SLIDE_W = 25.4
SLIDE_H = 14.29


def clear_slides(prs):
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        slide_id_list.remove(slide_id)


def set_text(shape, text, size=14, color=BLACK, bold=False, align=PP_ALIGN.CENTER):
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return text_frame


def add_title(slide, title, subtitle=None, section=None):
    if section:
        tag = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(SLIDE_W), Cm(0.18))
        tag.fill.solid()
        tag.fill.fore_color.rgb = section
        tag.line.color.rgb = section
    box = slide.shapes.add_textbox(Cm(0.75), Cm(0.42), Cm(24), Cm(0.9))
    set_text(box, title, size=24, color=ST_BLUE, bold=True, align=PP_ALIGN.LEFT)
    if subtitle:
        sub = slide.shapes.add_textbox(Cm(0.8), Cm(1.25), Cm(23.6), Cm(0.6))
        set_text(sub, subtitle, size=10.5, color=ST_GRAY, align=PP_ALIGN.LEFT)


def add_footer(slide, idx, total=24):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.75), Cm(13.25), Cm(23.9), Cm(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 225, 232)
    line.line.color.rgb = RGBColor(220, 225, 232)
    footer = slide.shapes.add_textbox(Cm(0.85), Cm(13.34), Cm(23.5), Cm(0.35))
    set_text(footer, "STM32Cube Preprocessing Pipeline | 40-minute technical presentation", size=7.5, color=ST_GRAY, align=PP_ALIGN.RIGHT)


def add_box(slide, x, y, w, h, text, fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12, bold=True, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    set_text(shape, text, size=font_size, color=BLACK, bold=bold, align=align)
    return shape


def add_label(slide, x, y, w, h, text, size=11, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    set_text(shape, text, size=size, color=color, bold=bold, align=align)
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=ST_BLUE, dashed=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(2)
    connector.line.end_arrowhead = True
    if dashed:
        connector.line.dash_style = 4
    return connector


def add_bullets(slide, x, y, w, h, title, bullets, color=ST_BLUE):
    add_label(slide, x, y, w, 0.55, title, size=15, color=color, bold=True)
    top = y + 0.78
    for bullet in bullets:
        add_label(slide, x + 0.15, top, 0.35, 0.35, "-", size=11.5, color=color, bold=True)
        add_label(slide, x + 0.55, top - 0.02, w - 0.55, 0.45, bullet, size=11, color=BLACK)
        top += 0.62


def add_bar(slide, x, y, label, value, color, max_value=100, width=11.5):
    add_label(slide, x, y - 0.05, 4.5, 0.4, label, size=10.5, bold=True)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + 4.9), Cm(y), Cm(width), Cm(0.38))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(232, 236, 242)
    bg.line.color.rgb = RGBColor(232, 236, 242)
    fg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + 4.9), Cm(y), Cm(width * value / max_value), Cm(0.38))
    fg.fill.solid()
    fg.fill.fore_color.rgb = color
    fg.line.color.rgb = color
    add_label(slide, x + 5.15 + width, y - 0.07, 1.8, 0.42, f"{value:g}", size=10.2, bold=True, color=color)


def add_code(slide, x, y, w, h, text, title=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(34, 39, 48)
    shape.line.color.rgb = RGBColor(34, 39, 48)
    if title:
        add_label(slide, x, y - 0.48, w, 0.35, title, size=10.5, color=ST_GRAY, bold=True)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Consolas"
    p.font.size = Pt(9.2)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT
    return shape


def add_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes


def load_template():
    last_error = None
    for candidate in TEMPLATE_CANDIDATES:
        if not candidate.exists():
            continue
        try:
            prs = Presentation(candidate)
            print(f"Using template: {candidate}")
            return prs
        except PermissionError as exc:
            print(f"Template locked, trying next: {candidate}")
            last_error = exc
        except Exception as exc:
            print(f"Template unavailable, trying next: {candidate} ({exc})")
            last_error = exc
    if last_error:
        raise RuntimeError("No readable PowerPoint template was available. Close PowerPoint/OneDrive preview for the deck and retry.") from last_error
    raise FileNotFoundError("No PowerPoint template candidate was found in docs/.")


def save_presentation(prs):
    try:
        prs.save(OUTPUT)
        return OUTPUT
    except PermissionError:
        fallback = ROOT / "STM32Cube_Preprocessing_Solution_40min_English_new.pptx"
        prs.save(fallback)
        return fallback


def new_slide(prs, idx, title, subtitle=None, section=ST_BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle, section=section)
    add_footer(slide, idx)
    return slide


def slide_01(prs):
    slide = new_slide(prs, 1, "STM32Cube Preprocessing Pipeline", "From raw GitHub evidence to a retrieval-ready knowledge base", section=ST_BLUE)
    add_label(slide, 1.0, 3.0, 23.2, 1.0, "40-minute technical presentation", size=28, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 3.2, 5.0, 5.0, 1.4, "GitHub data", fill=LIGHT_GRAY, line=ST_GRAY, font_size=14)
    add_arrow(slide, 8.4, 5.7, 9.5, 5.7)
    add_box(slide, 9.7, 5.0, 5.0, 1.4, "Python pipeline", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=14)
    add_arrow(slide, 14.9, 5.7, 16.0, 5.7)
    add_box(slide, 16.2, 5.0, 5.0, 1.4, "KB #793", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=14)
    add_label(slide, 2.0, 8.6, 21.4, 1.5, "Focus: the solution developed, the data structuring choices, the Python scripts, and the main technical challenges around Alfred, NovaX, and image processing.", size=14, color=BLACK, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 1 min. Open with the objective: this is not a generic RAG talk, it is about adapting preprocessing to the existing ST ChatGPT / KB platform.")


def slide_executive_summary(prs):
    slide = new_slide(prs, 2, "Executive Summary", "A preprocessing solution built to make ST GitHub knowledge usable by an industrial RAG", section=ST_BLUE)
    add_box(slide, 1.0, 2.2, 5.2, 1.45, "Problem\nRaw GitHub content is noisy, scattered, and weakly structured", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=12)
    add_arrow(slide, 6.4, 2.95, 7.4, 2.95, color=ST_ORANGE)
    add_box(slide, 7.6, 2.2, 5.2, 1.45, "Solution\nPython pipeline creates ST-ready structured artifacts", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_arrow(slide, 13.0, 2.95, 14.0, 2.95, color=ST_ORANGE)
    add_box(slide, 14.2, 2.2, 5.2, 1.45, "Platform\nArtifacts are adapted to ST ChatGPT RAG constraints", fill=RGBColor(235, 232, 255), line=PURPLE, font_size=12)
    add_arrow(slide, 19.6, 2.95, 20.6, 2.95, color=ST_ORANGE)
    add_box(slide, 20.8, 2.2, 3.5, 1.45, "Impact\nQI up to 94", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_bullets(slide, 1.2, 5.4, 10.8, 4.0, "What was delivered", ["End-to-end pipeline: ingestion, cleaning, enrichment, similarity, validation, delivery", "Resolver cases and diagnostic cards for support-oriented retrieval", "ALFRED Vision enrichment for issue images and PDF figures", "Upload workflow for KB #793 with schema gates"], color=ST_BLUE)
    add_bullets(slide, 13.0, 5.4, 10.8, 4.0, "What was evaluated", ["Each stage has quality metrics", "RAG configuration was tested: hybrid search, thresholds, retrieved docs, prompts", "Diagnostic cards + resolver cases produced the largest gain: +33 QI points", "Fallback routing was restricted to keep answers grounded in KB evidence"], color=ST_GREEN)
    add_notes(slide, "Timing: 2 min. Give the audience the full story before entering details.")


def slide_02(prs):
    slide = new_slide(prs, 2, "Talk Plan and Timing", "A compact 40-minute path: solution first, then technical challenges", section=ST_BLUE)
    items = [("1", "Business need", "4 min"), ("2", "Architecture", "5 min"), ("3", "Pipeline", "8 min"), ("4", "Preprocessing logic", "10 min"), ("5", "Evaluation + config", "8 min"), ("6", "Impact + demo", "5 min")]
    x = 1.2
    for n, title, timing in items:
        add_box(slide, x, 3.0, 3.5, 2.4, f"{n}\n{title}\n{timing}", fill=RGBColor(235, 240, 252), line=ST_BLUE, font_size=11)
        x += 3.9
    add_label(slide, 1.4, 7.2, 22.0, 1.2, "Short version: business need -> architecture -> pipeline -> intelligent artifacts -> evaluation -> operational use.", size=14, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 1 min. Tell the audience exactly where the presentation goes.")


def slide_03(prs):
    slide = new_slide(prs, 3, "Why This Solution Was Needed", "STM32Cube knowledge is rich, but scattered, noisy, and partly invisible to retrieval", section=ST_BLUE)
    add_bullets(slide, 1.2, 2.4, 10.5, 5.0, "Industrial support problem", ["STM32 engineers need fast access to known bugs, fixes, versions, and workarounds", "Knowledge is distributed across GitHub issues, PRs, commits, files, release notes, screenshots, and PDFs", "Manual search is slow and weakly traceable"])
    add_bullets(slide, 13.2, 2.4, 10.5, 5.0, "Pipeline objective", ["Convert raw GitHub evidence into KB-ready structured documents", "Preserve source traceability: repo, issue, URL, PR, commit", "Improve retrieval quality without redesigning the ST platform"])
    add_box(slide, 2.2, 9.8, 20.8, 1.45, "The project is a preprocessing and knowledge-structuring solution, not a supervised ML model.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=15)
    add_notes(slide, "Timing: 2 min. Anchor the project in the support need.")


def slide_04(prs):
    slide = new_slide(prs, 4, "Solution at a Glance", "A 7-stage Python pipeline that feeds KB #793 with structured technical evidence", section=ST_GREEN)
    stages = ["1. Ingestion", "2. Cleaning V1/V3", "3. Enrichment", "4. Similarity", "5. Chunking tests", "6. Validation", "7. Delivery"]
    x = 0.8
    for i, stage in enumerate(stages):
        add_box(slide, x, 3.0, 3.0, 1.15, stage, fill=ST_LIGHT_BLUE if i not in [1, 6] else RGBColor(225, 245, 235), line=ST_BLUE if i not in [1, 6] else ST_GREEN, font_size=10.5)
        if i < len(stages) - 1:
            add_arrow(slide, x + 3.05, 3.55, x + 3.55, 3.55)
        x += 3.45
    add_bullets(slide, 1.2, 5.8, 10.7, 4.0, "What the pipeline produces", ["Issues structured for troubleshooting", "Files cleaned and documented", "Resolver cases with Issue -> PR -> Commit evidence", "Diagnostic cards for recurring support cases"], color=ST_GREEN)
    add_bullets(slide, 13.2, 5.8, 10.7, 4.0, "Why it adds value", ["Improves signal density before embedding", "Keeps documents compatible with PARENT_CHILD chunking", "Makes images and PDF figures searchable", "Blocks malformed data before upload"], color=ST_BLUE)
    add_notes(slide, "Timing: 2 min. Explain the high-level flow and the four output artifact types.")


def slide_05(prs):
    slide = new_slide(prs, 5, "Existing Platform Constraints", "The pipeline was adapted to the ST retrieval platform; it did not redesign it", section=PURPLE)
    add_box(slide, 1.0, 2.2, 5.0, 1.3, "chunkStrategy\nPARENT_CHILD", fill=RGBColor(235, 232, 255), line=PURPLE, font_size=13)
    add_box(slide, 6.7, 2.2, 5.0, 1.3, "childChunkSize\n1000", fill=RGBColor(235, 232, 255), line=PURPLE, font_size=13)
    add_box(slide, 12.4, 2.2, 5.0, 1.3, "parentChunkSize\n3000", fill=RGBColor(235, 232, 255), line=PURPLE, font_size=13)
    add_box(slide, 18.1, 2.2, 5.0, 1.3, "embedding model\nAdaV2", fill=RGBColor(235, 232, 255), line=PURPLE, font_size=13)
    add_box(slide, 4.0, 5.2, 7.2, 1.25, "jobBatchSize = 60\nindexing throughput", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_box(slide, 14.0, 5.2, 7.2, 1.25, "embeddingsBatchSize = 20\nembedding API throughput", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_label(slide, 1.6, 8.2, 22.0, 1.2, "Design consequence: every exported document must be clean, compact, semantically dense, traceable, and stable under platform-side PARENT_CHILD chunking.", size=15, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 2 min. Stress that these were platform parameters, not project inventions.")


def slide_06(prs):
    slide = new_slide(prs, 6, "PARENT_CHILD Retrieval Flow", "Child chunks are indexed; parent chunks are returned as context", section=PURPLE)
    add_box(slide, 1.0, 2.5, 4.0, 1.2, "ST-ready document\nissue / file / card", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_arrow(slide, 5.2, 3.1, 6.4, 3.1)
    add_box(slide, 6.6, 1.8, 4.5, 2.6, "Parent chunk\n3000 chars\ncontext window unit", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_box(slide, 7.0, 2.2, 3.7, 0.55, "Child 1 - 1000 chars", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=9)
    add_box(slide, 7.0, 2.9, 3.7, 0.55, "Child 2 - 1000 chars", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=9)
    add_box(slide, 7.0, 3.6, 3.7, 0.55, "Child 3 - 1000 chars", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=9)
    add_arrow(slide, 11.4, 3.1, 13.2, 3.1)
    add_box(slide, 13.4, 2.5, 4.5, 1.2, "Vector index\nAdaV2", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_box(slide, 18.9, 2.5, 4.2, 1.2, "User query\nmatched to child", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_arrow(slide, 18.8, 3.1, 18.0, 3.1, color=ST_ORANGE)
    add_arrow(slide, 13.6, 4.5, 11.2, 6.2, color=ST_ORANGE, dashed=True)
    add_box(slide, 6.6, 6.1, 4.5, 1.25, "Parent returned\nto GPT-5.1", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_arrow(slide, 11.3, 6.7, 13.1, 6.7, color=ST_ORANGE)
    add_box(slide, 13.3, 6.1, 4.7, 1.25, "Grounded answer\nwith citations", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_label(slide, 1.0, 9.4, 22.8, 1.4, "Pipeline implication: each document must stay coherent when split into 1000-char child chunks and 3000-char parent context blocks.", size=14, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 2 min. This is the central platform adaptation slide.")


def slide_07(prs):
    slide = new_slide(prs, 7, "How GitHub Data Is Structured", "Raw GitHub material is transformed into four complementary knowledge artifacts", section=ST_GREEN)
    artifacts = [("Structured issues", "Symptom, root cause hint, metadata, source URL"), ("Documented files", "HAL/BSP/CMSIS/README/PDF text, cleaned by type"), ("Resolver cases", "Issue -> PR -> Commit proof chain"), ("Diagnostic cards", "Symptom -> root cause -> fix -> prevention + aliases")]
    y = 2.4
    colors = [ST_BLUE, CYAN, ST_ORANGE, ST_GREEN]
    for (title, desc), color in zip(artifacts, colors):
        add_box(slide, 1.0, y, 5.0, 1.0, title, fill=RGBColor(235, 240, 252), line=color, font_size=12)
        add_label(slide, 6.4, y + 0.18, 17.0, 0.5, desc, size=12.5)
        y += 1.7
    add_box(slide, 2.0, 10.2, 20.8, 1.15, "The artifacts are not generic chunks: they are retrieval-oriented support evidence units.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=15)
    add_notes(slide, "Timing: 2 min. Explain why four artifact types are needed.")


def slide_08(prs):
    slide = new_slide(prs, 8, "Python Scripts Developed", "Each stage is isolated, replayable, and produces inspectable JSON artifacts", section=ST_GREEN)
    rows = [("Ingestion", "fetch_issues.py, fetch_files.py, fetch_prs.py, fetch_commits.py"), ("Cleaning", "clean_issues.py, clean_files.py, clean_files_v3.py"), ("Enrichment", "issues_to_docs_v2.py, files_to_docs_v2.py"), ("Similarity", "compute_issue_similarity_v2.py"), ("Delivery", "export_st_ready_issues.py, export_st_ready_files.py, resolver/diagnostic exporters"), ("Automation", "run_full_workflow.py, upload scripts, Alfred enrichment scripts")]
    y = 2.2
    for stage, scripts in rows:
        add_box(slide, 1.0, y, 4.2, 0.8, stage, fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=11.5)
        add_label(slide, 5.6, y + 0.12, 18.3, 0.5, scripts, size=10.8)
        y += 1.05
    add_code(slide, 1.2, 9.4, 22.0, 1.4, "python -m pipeline.run_full_workflow --export-st-ready --skip-chunking --continue-on-error", title="Main execution command")
    add_notes(slide, "Timing: 2 min. Show that the project is an implemented pipeline, not just a concept.")


def slide_09(prs):
    slide = new_slide(prs, 9, "Ingestion: Capturing GitHub Evidence", "Data Understanding becomes a controlled collection process", section=ST_GREEN)
    add_bullets(slide, 1.1, 2.4, 10.7, 4.5, "Collected inputs", ["Issues and comments", "Pull requests and reviews", "Commits and links", "Source files, README, release notes, PDFs", "Image URLs from comments and issue bodies"], color=ST_BLUE)
    add_bullets(slide, 13.0, 2.4, 10.7, 4.5, "Engineering constraints", ["GitHub pagination and rate limits", "Multiple repositories per series", "Missing raw files in cached runs", "Need for source URLs and traceability"], color=ST_ORANGE)
    add_box(slide, 2.0, 9.3, 20.8, 1.4, "Output: raw_issues_*.json, raw_files_*.json, raw_prs_*.json, raw_commits_*.json", fill=LIGHT_GRAY, line=ST_GRAY, font_size=13)
    add_notes(slide, "Timing: 1.5 min.")


def slide_10(prs):
    slide = new_slide(prs, 10, "Cleaning V3: Making Content Embedding-Friendly", "The goal is not only to remove noise, but to improve semantic indexing quality", section=ST_GREEN)
    add_box(slide, 1.0, 2.2, 4.2, 1.0, "Release Notes\nHTML", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_box(slide, 6.0, 2.2, 4.2, 1.0, "README\nMarkdown", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_box(slide, 11.0, 2.2, 4.2, 1.0, "Source code\n.c / .h", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_box(slide, 16.0, 2.2, 4.2, 1.0, "CMSIS header\n>100 KB", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    for x, txt in [(1.0, "strip_html()\n-33%"), (6.0, "clean_markdown()"), (11.0, "strip_license()"), (16.0, "filter_cmsis()\n-96%")]:
        add_arrow(slide, x + 2.1, 3.35, x + 2.1, 4.4)
        add_box(slide, x, 4.55, 4.2, 1.0, txt, fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=11)
    add_bar(slide, 2.0, 7.5, "CMSIS reduction", 96.4, ST_GREEN, width=11.0)
    add_bar(slide, 2.0, 8.7, "Corpus H7 reduction", 38.0, ST_ORANGE, width=11.0)
    add_bar(slide, 2.0, 9.9, "Signal density gain", 31.0, CYAN, width=11.0)
    add_notes(slide, "Timing: 2 min. Tie cleaning directly to AdaV2 embedding quality.")


def slide_11(prs):
    slide = new_slide(prs, 11, "Enrichment: From Text to Searchable Metadata", "Heuristics add structure without supervised training data", section=ST_GREEN)
    add_bullets(slide, 1.1, 2.3, 10.6, 5.5, "Fields added", ["layer: HAL, BSP, CMSIS, Middleware", "component: SPI, I2C, DMA, USB, etc.", "board: NUCLEO, Discovery, Eval", "severity and issue kind", "evidence strength and validity flags"], color=ST_BLUE)
    add_bullets(slide, 13.0, 2.3, 10.6, 5.5, "Why it matters", ["Faceted filtering and re-ranking", "Better query matching under hybrid search", "Cleaner answers with domain anchors", "Avoids pretending that a supervised ML model exists"], color=ST_GREEN)
    add_bar(slide, 2.2, 9.2, "Layer", 88.2, ST_GREEN, width=11.0)
    add_bar(slide, 2.2, 10.15, "Component", 72.6, ST_ORANGE, width=11.0)
    add_bar(slide, 2.2, 11.1, "Board", 45.3, ST_RED, width=11.0)
    add_notes(slide, "Timing: 2 min. Explain board coverage as source-data limitation.")


def slide_12(prs):
    slide = new_slide(prs, 12, "Similarity and Linking", "TF-IDF cosine links related issues without training a model", section=ST_GREEN)
    add_box(slide, 1.2, 2.2, 5.2, 1.15, "Issue text\ntitle + body + comments", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_arrow(slide, 6.7, 2.8, 8.0, 2.8)
    add_box(slide, 8.2, 2.2, 4.5, 1.15, "TF-IDF\nvectorization", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=11)
    add_arrow(slide, 12.9, 2.8, 14.2, 2.8)
    add_box(slide, 14.4, 2.2, 4.5, 1.15, "Cosine\nsimilarity", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=11)
    add_arrow(slide, 19.1, 2.8, 20.4, 2.8)
    add_box(slide, 20.6, 2.2, 3.5, 1.15, "related\nissues", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=11)
    add_label(slide, 1.5, 5.5, 22.0, 1.1, "Threshold: cosine similarity >= 0.3. Strict enough to avoid noisy links; H7 coverage is 22%, reflecting issue diversity rather than a pipeline failure.", size=14, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_code(slide, 2.0, 8.0, 20.8, 1.7, "sim(d_i, d_j) = (v_i . v_j) / (||v_i|| ||v_j||)\nrelated(d_i) = { d_j | sim(d_i, d_j) >= 0.3 }", title="Simplified formula")
    add_notes(slide, "Timing: 1.5 min.")


def slide_13(prs):
    slide = new_slide(prs, 13, "Delivery: KB-Ready Outputs", "The final stage converts enriched documents into platform-ready JSON", section=ST_GREEN)
    add_box(slide, 1.0, 2.3, 5.2, 1.25, "Issues JSON\nproblem evidence", fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=12)
    add_box(slide, 7.0, 2.3, 5.2, 1.25, "Files JSON\nclean technical content", fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=12)
    add_box(slide, 13.0, 2.3, 5.2, 1.25, "Resolver cases\nIssue -> PR -> Commit", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_box(slide, 19.0, 2.3, 5.2, 1.25, "Diagnostic cards\ntroubleshooting logic", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_bullets(slide, 1.2, 5.3, 10.5, 4.2, "Compatibility with KB #793", ["One datasource per series and artifact type", "JSON Schema validation gates", "Upload scripts reuse datasource IDs", "Delete+Add workflow avoids naming conflicts"], color=ST_BLUE)
    add_bullets(slide, 13.0, 5.3, 10.5, 4.2, "Retrieval value", ["More entry points for user formulations", "Evidence chains for traceability", "Structured support-oriented answers", "Less dependence on raw text chunks"], color=ST_GREEN)
    add_notes(slide, "Timing: 2 min.")


def slide_resolver_construction(prs):
    slide = new_slide(prs, 14, "Resolver Cases: From GitHub Discussion to Proof Chain", "Easy idea: keep the path that proves a fix, not only the discussion around it", section=ST_GREEN)
    add_box(slide, 1.0, 2.4, 4.2, 1.1, "Issue\nproblem reported", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_arrow(slide, 5.4, 3.0, 6.4, 3.0, color=ST_GREEN)
    add_box(slide, 6.6, 2.4, 4.2, 1.1, "Pull Request\nfix proposed", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_arrow(slide, 11.0, 3.0, 12.0, 3.0, color=ST_GREEN)
    add_box(slide, 12.2, 2.4, 4.2, 1.1, "Commit\ncode changed", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_arrow(slide, 16.6, 3.0, 17.6, 3.0, color=ST_GREEN)
    add_box(slide, 17.8, 2.4, 5.6, 1.1, "Resolver case\nKB-ready proof", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_bullets(slide, 1.2, 5.2, 10.5, 4.0, "How it is built", ["Collect issue -> PR -> commit links", "Classify resolution status: merged_fix, candidate_fix, reference_only", "Extract changed files and impacted modules", "Write a compact resolver_card_text"], color=ST_GREEN)
    add_bullets(slide, 13.0, 5.2, 10.5, 4.0, "Added value", ["The answer can cite a concrete fix path", "The engineer sees what changed, where, and why", "The LLM gets source-grounded evidence", "Traceability improves trust in the response"], color=ST_BLUE)
    add_box(slide, 2.0, 10.4, 20.8, 1.1, "Simple explanation: a resolver case is the short technical story of how a reported problem was resolved in code.", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=14)
    add_notes(slide, "Timing: 2 min. Explain resolver cases as proof chains, not complex data structures.")


def slide_resolver_example(prs):
    slide = new_slide(prs, 15, "Concrete Resolver Case Example", "A compact evidence card that the RAG can retrieve and cite", section=ST_GREEN)
    add_code(slide, 1.0, 2.1, 10.8, 8.6, "{\n  \"issue_number\": 42,\n  \"repo\": \"stm32cubeh7\",\n  \"resolution_status\": \"merged_fix\",\n  \"link_confidence\": \"high\",\n  \"primary_pr_number\": 87,\n  \"primary_changed_file\": \"Src/stm32h7xx_hal_spi.c\",\n  \"impacted_modules\": [\"SPI\", \"DMA\"],\n  \"resolver_card_text\": \"PROBLEM: SPI DMA hard fault... FIX: cache maintenance added... PROOF: PR #87 merged...\"\n}", title="ST-ready resolver case")
    add_bullets(slide, 13.0, 2.3, 10.5, 3.4, "Why this helps retrieval", ["Keywords: SPI, DMA, hard fault, cache", "Exact source: issue number, PR number, changed file", "Actionable content: problem + fix + proof"], color=ST_BLUE)
    add_bullets(slide, 13.0, 7.0, 10.5, 3.0, "What the answer can say", ["This issue was fixed in PR #87", "The modified file was stm32h7xx_hal_spi.c", "The root cause was missing cache maintenance"], color=ST_GREEN)
    add_notes(slide, "Timing: 1.5 min. Use this slide to make resolver cases tangible.")


def slide_diagnostic_construction(prs):
    slide = new_slide(prs, 16, "Diagnostic Cards: Turning Issues into Troubleshooting Logic", "Easy idea: write the way a support engineer reasons", section=ST_GREEN)
    steps = [("Symptom", "What the user observes"), ("Hypotheses", "What could cause it"), ("Validation", "How to check it"), ("Root cause", "What really explains it"), ("Fix", "What to do"), ("Prevention", "How to avoid it")]
    x = 0.8
    for title, desc in steps:
        add_box(slide, x, 2.6, 3.55, 1.6, f"{title}\n{desc}", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=10.3)
        if title != "Prevention":
            add_arrow(slide, x + 3.6, 3.4, x + 4.0, 3.4, color=ST_GREEN)
        x += 4.0
    add_bullets(slide, 1.2, 5.6, 10.8, 4.0, "Construction rule", ["Select issues with strong technical evidence", "Compress long discussions into a clean reasoning path", "Add query aliases for user reformulations", "Keep GitHub source metadata for traceability"], color=ST_BLUE)
    add_bullets(slide, 13.0, 5.6, 10.8, 4.0, "Why it is powerful", ["A support answer is not just a document match", "It needs diagnosis, verification, and fix", "The card gives the LLM that structure directly", "This produced the largest QI jump: +33 points"], color=ST_GREEN)
    add_notes(slide, "Timing: 2 min. Explain diagnostic cards as support reasoning templates.")


def slide_diagnostic_example(prs):
    slide = new_slide(prs, 17, "Concrete Diagnostic Card Example", "A card adds aliases so the same problem can be found with different wording", section=ST_GREEN)
    add_code(slide, 1.0, 2.0, 10.8, 8.9, "{\n  \"diagnostic_id\": \"DIAG-H7-042\",\n  \"symptom\": \"Hard fault after second SPI DMA transfer\",\n  \"root_cause\": \"D-cache not invalidated before DMA reception\",\n  \"resolution\": \"Invalidate D-cache or use non-cacheable SRAM\",\n  \"query_aliases\": [\n    \"SPI DMA hard fault STM32H7\",\n    \"cache coherency SPI circular mode\",\n    \"HAL_SPI_TransmitReceive_DMA crash\"\n  ],\n  \"evidence_strength\": \"high\"\n}", title="ST-ready diagnostic card")
    add_box(slide, 13.0, 2.3, 9.8, 1.4, "User asks:\n\"My SPI crashes after DMA transfer\"", fill=LIGHT_GRAY, line=ST_GRAY, font_size=13)
    add_arrow(slide, 17.9, 3.9, 17.9, 5.0, color=ST_ORANGE)
    add_box(slide, 13.0, 5.1, 9.8, 1.5, "Hybrid search matches aliases\n+ technical anchors", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_arrow(slide, 17.9, 6.8, 17.9, 7.9, color=ST_GREEN)
    add_box(slide, 13.0, 8.0, 9.8, 1.6, "Answer includes symptom, root cause, fix, and source evidence", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_notes(slide, "Timing: 1.5 min. Show aliases as the bridge between human wording and KB wording.")


def slide_resolver_diagnostic_combined(prs):
    slide = new_slide(prs, 18, "Resolver Cases + Diagnostic Cards", "The added intelligence layer between raw GitHub data and ST ChatGPT retrieval", section=ST_GREEN)
    add_box(slide, 1.0, 2.2, 10.6, 2.6, "Resolver Case\n\nIssue -> PR -> Commit -> Changed file\n\nPurpose: prove how a reported problem was fixed in code.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_box(slide, 13.0, 2.2, 10.6, 2.6, "Diagnostic Card\n\nSymptom -> Hypotheses -> Validation -> Root cause -> Fix\n\nPurpose: make support reasoning directly retrievable.", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_box(slide, 1.4, 6.0, 3.4, 1.0, "Raw issue", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11.5)
    add_arrow(slide, 5.0, 6.5, 6.2, 6.5, color=ST_ORANGE)
    add_box(slide, 6.4, 6.0, 4.3, 1.0, "Evidence chain", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=11.5)
    add_arrow(slide, 10.9, 6.5, 12.1, 6.5, color=ST_GREEN)
    add_box(slide, 12.3, 6.0, 4.4, 1.0, "Reasoning card", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=11.5)
    add_arrow(slide, 16.9, 6.5, 18.1, 6.5, color=ST_BLUE)
    add_box(slide, 18.3, 6.0, 4.8, 1.0, "Better answer", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=11.5)
    add_bullets(slide, 1.2, 8.1, 10.8, 2.4, "Simple example", ["Question: 'My SPI DMA crashes on STM32H7'", "Resolver: points to PR/commit/changed HAL file", "Diagnostic card: explains cache coherency, validation, and fix"], color=ST_BLUE)
    add_bullets(slide, 13.0, 8.1, 10.8, 2.4, "Measured value", ["Largest improvement in the evaluation loop", "+33 Quality Index points", "Answers become more traceable and more actionable"], color=ST_GREEN)
    add_notes(slide, "Timing: 2.5 min. This replaces separate detailed slides when time is limited.")


def slide_artifact_added_value(prs):
    slide = new_slide(prs, 18, "Why Resolver Cases and Diagnostic Cards Changed the Game", "The main gain came from structure, not from a bigger model", section=ST_GREEN)
    add_box(slide, 1.0, 2.3, 6.8, 2.3, "Raw issue\nLong thread, mixed context, comments, screenshots, partial fixes", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=13)
    add_arrow(slide, 8.0, 3.5, 9.2, 3.5, color=ST_ORANGE)
    add_box(slide, 9.4, 2.3, 6.8, 2.3, "Resolver case\nProof chain: issue -> PR -> commit -> changed file", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_arrow(slide, 16.4, 3.5, 17.6, 3.5, color=ST_GREEN)
    add_box(slide, 17.8, 2.3, 6.4, 2.3, "Diagnostic card\nReasoning path: symptom -> root cause -> fix", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_bar(slide, 2.0, 7.2, "Issues only baseline", 54, ST_RED, width=12.0)
    add_bar(slide, 2.0, 8.5, "After prompt + hybrid", 61, ST_ORANGE, width=12.0)
    add_bar(slide, 2.0, 9.8, "With diagnostic cards + resolver", 94, ST_GREEN, width=12.0)
    add_label(slide, 2.0, 11.3, 21.0, 0.8, "Takeaway: the best improvement was making the knowledge easier to retrieve and easier for the LLM to reason over.", size=13, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 2 min. Use the QI jump to make added value obvious.")


def slide_stage_evaluation(prs):
    slide = new_slide(prs, 19, "Each Pipeline Stage Was Evaluated", "Evaluation is not only the final chatbot score; every stage has quality gates", section=CYAN)
    rows = [("Ingestion", "Completeness", "99.7% body / 100% title"), ("Cleaning V3", "Noise reduction", "-96.4% CMSIS / -38% corpus"), ("Enrichment", "Metadata coverage", "Layer 88.2%, component 72.6%"), ("Alfred", "Contextual trust", "Issue images 98.5%, PDF figures 41.7%"), ("Similarity", "Related issue coverage", "22% on H7"), ("Delivery", "Schema validity", "100% eligible exports")]
    y = 2.2
    for stage, metric, result in rows:
        add_box(slide, 1.0, y, 4.1, 0.72, stage, fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=10.5)
        add_box(slide, 5.5, y, 5.2, 0.72, metric, fill=LIGHT_GRAY, line=ST_GRAY, font_size=10.5)
        add_label(slide, 11.2, y + 0.1, 12.0, 0.42, result, size=10.5)
        y += 0.92
    add_box(slide, 2.0, 9.7, 20.8, 1.2, "This evaluation helped distinguish pipeline defects from source-data limitations, for example the low board coverage caused by missing board mentions in issues.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_notes(slide, "Timing: 2 min. Show that every stage had a metric.")


def slide_chatgpt_config_knobs(prs):
    slide = new_slide(prs, 20, "ST ChatGPT / RAG Configuration I Could Tune", "The pipeline covered not only data preparation but also platform configuration experiments", section=PURPLE)
    knobs = [("Search type", "Semantic / Full-text / Hybrid", "Hybrid retained"), ("Semantic threshold", "0.45 / 0.55 / 0.65", "S0=0.55, S1=0.45"), ("Semantic doc count", "18 / 25 / 35", "25 stable, 35 recall+"), ("Full-text doc count", "8 / 12 / 15 / 20", "12 stable"), ("System prompt", "KB-first behavior", "+3 QI points"), ("Tool description", "NovaX fallback strictness", "Prevents KB bypass")]
    y = 2.1
    for name, tested, decision in knobs:
        add_box(slide, 1.0, y, 4.3, 0.75, name, fill=RGBColor(235, 232, 255), line=PURPLE, font_size=10.5)
        add_label(slide, 5.7, y + 0.1, 8.0, 0.42, tested, size=10.5)
        add_label(slide, 14.4, y + 0.1, 8.6, 0.42, decision, size=10.5, color=ST_BLUE, bold=True)
        y += 0.95
    add_box(slide, 2.0, 9.4, 20.8, 1.35, "Important: these are platform configuration parameters. The project did not redesign the RAG architecture; it adapted preprocessing and tested safe settings around it.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_notes(slide, "Timing: 2 min. Clarify what was configurable and what was imposed.")


def slide_config_results(prs):
    slide = new_slide(prs, 21, "Configuration Experiments: What Worked", "Hybrid retrieval was necessary because STM32 questions mix natural language and exact API names", section=PURPLE)
    add_bar(slide, 2.0, 2.8, "S1 Recall+ Hybrid", 66, ST_GREEN, width=12.0)
    add_bar(slide, 2.0, 4.0, "S0 Balanced Hybrid", 65, ST_GREEN, width=12.0)
    add_bar(slide, 2.0, 5.2, "S2 Precision+ Hybrid", 63, ST_ORANGE, width=12.0)
    add_bar(slide, 2.0, 6.4, "S3 Semantic-heavy", 51, ST_RED, width=12.0)
    add_bar(slide, 2.0, 7.6, "S4 Fulltext-heavy", 47, ST_RED, width=12.0)
    add_bar(slide, 2.0, 8.8, "S5 Semantic-only", 41, ST_RED, width=12.0)
    add_bullets(slide, 16.0, 3.0, 7.5, 4.4, "Interpretation", ["Hybrid is the best compromise", "Too much semantic loses API anchors", "Too much full-text loses reformulations", "Recall+ gives only +1, so noise must be controlled"], color=PURPLE)
    add_notes(slide, "Timing: 2 min. This slide links configuration choices to measured QI.")


def slide_system_prompt_and_routing(prs):
    slide = new_slide(prs, 22, "System Prompt and Tool Descriptions", "Small text changes can strongly affect routing behavior", section=PURPLE)
    add_box(slide, 1.0, 2.2, 10.4, 2.5, "System prompt\n\n- Always search KB first\n- Cite GitHub evidence\n- Prefer concrete fixes\n- Avoid unsupported generic answers", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_box(slide, 13.0, 2.2, 10.4, 2.5, "NovaX tool description\n\n- Use only if KB has no relevant answer\n- Forward the exact question\n- Do not bypass KB evidence", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_box(slide, 2.0, 6.4, 9.8, 2.2, "Before\nFallback too permissive\nQuality drops because answers become less traceable", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=13)
    add_arrow(slide, 12.0, 7.5, 13.1, 7.5, color=ST_ORANGE)
    add_box(slide, 13.4, 6.4, 9.8, 2.2, "After\nKB-first routing\nAnswers stay grounded in issues, files, resolver cases and diagnostic cards", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_notes(slide, "Timing: 2 min. Explain that prompt engineering here is operational routing control.")


def slide_14(prs):
    slide = new_slide(prs, 14, "Challenge 1: Calling Other Personas", "Fallback is useful only if it does not bypass the KB", section=ST_ORANGE)
    add_box(slide, 1.0, 2.1, 4.0, 1.2, "User question", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_arrow(slide, 5.2, 2.7, 6.7, 2.7)
    add_box(slide, 6.9, 1.8, 5.0, 1.8, "ST GitHub Analyzer\nmain persona", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_arrow(slide, 12.1, 2.3, 15.0, 1.7, color=ST_GREEN)
    add_box(slide, 15.2, 1.0, 5.0, 1.4, "KB #793\nALWAYS FIRST", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_arrow(slide, 12.1, 3.0, 15.0, 4.7, color=ST_ORANGE, dashed=True)
    add_box(slide, 15.2, 4.0, 5.0, 1.4, "NovaX\nONLY IF KB EMPTY", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_box(slide, 2.0, 8.0, 9.8, 2.2, "Problem found\nA permissive fallback description made the persona call the fallback too early, losing GitHub citations and lowering answer traceability.", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=12)
    add_box(slide, 13.2, 8.0, 9.8, 2.2, "Fix implemented\nRestrictive tool description: search the KB first; call NovaX only when no relevant KB evidence exists.", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_notes(slide, "Timing: 2 min. Explain routing as a quality parameter.")


def slide_15(prs):
    slide = new_slide(prs, 15, "Challenge 2: Image Content Is Invisible to Text Search", "Screenshots and PDF figures carry support evidence that embeddings cannot see directly", section=ST_ORANGE)
    add_box(slide, 1.2, 2.2, 9.8, 2.4, "Before\n[IMAGE ATTACHED]\nThe KB sees only a placeholder, not the error dialog, IDE log, waveform, or configuration screen.", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=12)
    add_arrow(slide, 11.3, 3.4, 13.2, 3.4, color=ST_ORANGE)
    add_box(slide, 13.5, 2.2, 9.8, 2.4, "After Alfred Vision\n[IMAGE DESCRIPTION]\nThe screenshot becomes searchable text with STM32-specific diagnostic anchors.", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_bullets(slide, 1.4, 6.5, 10.8, 3.4, "Issue images", ["Usually grounded in issue text", "Often include IDE logs, error strings, config screens", "High contextual trust: 98.5%"], color=ST_GREEN)
    add_bullets(slide, 13.2, 6.5, 10.8, 3.4, "PDF figures", ["Useful for recall but harder to validate", "Page text may not repeat figure labels", "Contextual trust: 41.7%"], color=ST_ORANGE)
    add_notes(slide, "Timing: 2 min. Make the distinction between API success and contextual confidence.")


def slide_16(prs):
    slide = new_slide(prs, 16, "Alfred Vision Pipeline", "From image URLs and PDF crops to searchable descriptions", section=ST_ORANGE)
    nodes = [("Image marker\n[IMAGE ATTACHED]", 1.0), ("URL / crop\nextraction", 5.6), ("Alfred Vision\nAPI call", 10.2), ("Description\ntext", 14.8), ("ST-ready\ninjection", 19.4)]
    for text, x in nodes:
        add_box(slide, x, 3.0, 3.6, 1.2, text, fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=11)
    for x in [4.7, 9.3, 13.9, 18.5]:
        add_arrow(slide, x, 3.6, x + 0.7, 3.6, color=ST_ORANGE)
    add_code(slide, 1.2, 6.3, 22.0, 1.7, "python pipeline_Automation/alfred/enrich_json_images_with_alfred.py --input <issues_with_images.json> --inplace\npython pipeline_Automation/alfred/enrich_pdf_figures_with_alfred.py --repo STM32CubeU5", title="Representative commands")
    add_bullets(slide, 1.4, 9.3, 21.0, 2.2, "Technical safeguards", ["Avoid duplicate image markers in root cause summaries", "Track image_urls to prevent URL exhaustion", "Evaluate contextual trust after enrichment"], color=ST_BLUE)
    add_notes(slide, "Timing: 2 min.")


def slide_17(prs):
    slide = new_slide(prs, 17, "Alfred Evaluation Results", "Measured confidence, not just API success", section=ST_ORANGE)
    add_bar(slide, 2.0, 3.0, "Issue images contextual trust", 98.5, ST_GREEN, width=12.5)
    add_bar(slide, 2.0, 4.3, "PDF figures contextual trust", 41.7, ST_ORANGE, width=12.5)
    add_bar(slide, 2.0, 5.6, "Total Alfred contextual trust", 67.4, CYAN, width=12.5)
    add_box(slide, 2.0, 8.0, 9.8, 2.0, "Issue images\n200 evaluated\n3 low coherence", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_box(slide, 13.3, 8.0, 9.8, 2.0, "PDF figures\n242 evaluated\n141 low coherence", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_label(slide, 2.0, 11.3, 21.0, 0.8, "Interpretation: PDF figure descriptions are kept for multimodal recall, but the report clearly separates them from the mature issue-image pipeline.", size=12.5, color=ST_GRAY, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 1.5 min.")


def slide_18(prs):
    slide = new_slide(prs, 18, "KB Upload and Indexing Workflow", "Delivery is part of the solution, not an afterthought", section=CYAN)
    add_box(slide, 1.0, 2.4, 4.2, 1.2, "ST-ready JSON", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_arrow(slide, 5.4, 3.0, 6.6, 3.0, color=CYAN)
    add_box(slide, 6.8, 2.4, 4.2, 1.2, "Schema gate", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_arrow(slide, 11.2, 3.0, 12.4, 3.0, color=CYAN)
    add_box(slide, 12.6, 2.4, 4.2, 1.2, "Upload script", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_arrow(slide, 17.0, 3.0, 18.2, 3.0, color=CYAN)
    add_box(slide, 18.4, 2.4, 4.2, 1.2, "KB #793", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_code(slide, 1.2, 5.5, 22.0, 2.2, "python pipeline_Automation/upload/Add_Data_Source_Files.py --kb 793 --operation add \\\n  --datasource-id <ID> --processor JSON --processor-params base64:<b64> \\\n  --json-split-size 1000 --files <file.json>", title="Upload command shape")
    add_bullets(slide, 1.4, 9.2, 21.0, 2.0, "Why upload engineering matters", ["Stable datasource IDs per series and artifact type", "Base64 processor params avoid PowerShell parsing problems", "Delete+Add avoids duplicate filename conflicts", "Validation gates prevent corrupted data from reaching the KB"], color=CYAN)
    add_notes(slide, "Timing: 2 min.")


def slide_19(prs):
    slide = new_slide(prs, 19, "Automation and Operational Robustness", "The pipeline must survive partial data, missing linked repositories, and repeated runs", section=CYAN)
    add_bullets(slide, 1.1, 2.4, 10.8, 5.0, "Robustness fixes", ["Skip missing raw issue files instead of crashing", "Skip unavailable delivery inputs with warnings", "Scope Alfred enrichment to the active series", "Refresh file/PDF detection without full Git sync"], color=ST_BLUE)
    add_bullets(slide, 13.0, 2.4, 10.8, 5.0, "Automation design", ["run_full_workflow orchestrates the stages", "--continue-on-error keeps other series running", "--skip-ingestion reuses cached artifacts", "Schema validation gates before upload"], color=CYAN)
    add_box(slide, 2.0, 9.6, 20.8, 1.2, "Operational lesson: automation quality is not only successful runs; it is predictable behavior when a repo or artifact is missing.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=14)
    add_notes(slide, "Timing: 2 min.")


def slide_20(prs):
    slide = new_slide(prs, 20, "Results and Measured Impact", "The system improves answer quality by adding structured, traceable evidence", section=CYAN)
    add_bar(slide, 2.0, 2.8, "Alfred baseline score", 58.0, ST_RED, width=12.0)
    add_bar(slide, 2.0, 4.1, "Persona KB score", 80.5, ST_GREEN, width=12.0)
    add_bar(slide, 2.0, 5.4, "H7 Quality Index with Diagnostic Cards", 94.0, ST_GREEN, width=12.0)
    add_box(slide, 2.0, 8.0, 6.2, 1.7, "+39%\nquality gain vs generic Alfred", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=14)
    add_box(slide, 9.5, 8.0, 6.2, 1.7, "+33 pts\nDiagnostic Cards impact", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=14)
    add_box(slide, 17.0, 8.0, 6.2, 1.7, "87%\nverifiable citations", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=14)
    add_notes(slide, "Timing: 1.5 min. Keep this short and evidence-focused.")


def slide_21(prs):
    slide = new_slide(prs, 21, "What Was Hardest Technically", "Three recurring problems shaped the final architecture", section=ST_ORANGE)
    add_box(slide, 1.2, 2.6, 6.8, 3.0, "1. Data heterogeneity\nHTML, Markdown, C headers, PDFs, issues, comments, screenshots", fill=LIGHT_GRAY, line=ST_GRAY, font_size=13)
    add_box(slide, 9.3, 2.6, 6.8, 3.0, "2. Agent routing\nFallback can reduce quality if it bypasses the KB", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=13)
    add_box(slide, 17.4, 2.6, 6.8, 3.0, "3. Image evidence\nVisual content must become text and must be validated", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=13)
    add_label(slide, 2.0, 8.2, 21.0, 1.2, "The final pipeline is the result of these constraints: not only extraction, but robust structuring, validation, routing, and enrichment.", size=15, color=ST_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Timing: 2 min. This is the transition to the conclusion.")


def slide_22(prs):
    slide = new_slide(prs, 22, "Key Takeaways", "What the audience should remember", section=ST_BLUE)
    add_box(slide, 1.5, 2.4, 22.0, 1.2, "The platform architecture was given; the pipeline was designed to make high-quality, traceable knowledge fit that architecture.", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=15)
    add_box(slide, 1.5, 4.4, 22.0, 1.2, "Structured artifacts are the main value: issues, files, resolver cases, and diagnostic cards are retrieval-ready support evidence.", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=15)
    add_box(slide, 1.5, 6.4, 22.0, 1.2, "The hardest parts were not only NLP: image enrichment, persona routing, schema validation, and KB upload stability were equally critical.", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=15)
    add_box(slide, 1.5, 8.4, 22.0, 1.2, "The result is a reproducible, multi-series, KB-ready preprocessing workflow for STM32Cube technical support.", fill=LIGHT_GRAY, line=ST_GRAY, font_size=15)
    add_notes(slide, "Timing: 1 min. Summarize with confidence.")


def slide_23(prs):
    slide = new_slide(prs, 23, "Backup: Representative JSON Artifact", "Example of a diagnostic card produced by the pipeline", section=ST_GRAY)
    add_code(slide, 1.0, 2.2, 23.0, 8.8, "{\n  \"diagnostic_id\": \"DIAG-H7-042\",\n  \"repo\": \"stm32cubeh7\",\n  \"symptom\": \"Hard fault after second SPI DMA transfer\",\n  \"root_cause\": \"D-cache not invalidated before DMA reception\",\n  \"resolution\": \"Call SCB_InvalidateDCache_by_Addr() or use non-cacheable SRAM\",\n  \"query_aliases\": [\n    \"SPI DMA hard fault STM32H7\",\n    \"cache coherency SPI circular mode\",\n    \"HAL_SPI_TransmitReceive_DMA crash\"\n  ],\n  \"evidence_strength\": \"high\"\n}", title="Diagnostic card excerpt")
    add_notes(slide, "Backup. Use if asked what structured output actually looks like.")


def slide_24(prs):
    slide = new_slide(prs, 24, "Backup: Main Commands", "Operational entry points used during development and validation", section=ST_GRAY)
    add_code(slide, 1.0, 2.2, 23.0, 7.8, "# Full pipeline with ST-ready export\npython -m pipeline.run_full_workflow --export-st-ready --skip-chunking --continue-on-error\n\n# Schema validation gate\npython -m pipeline.evaluation.validate_schemas\n\n# Alfred image enrichment\npython pipeline_Automation/alfred/enrich_json_images_with_alfred.py --input <file> --inplace\n\n# Upload to KB #793\npython pipeline_Automation/upload/Add_Data_Source_Files.py --kb 793 --operation add --datasource-id <ID> --processor JSON --files <file.json>", title="Commands")
    add_notes(slide, "Backup. Use if asked how the pipeline is executed.")


SLIDE_BUILDERS = [
    # 1. Executive Summary
    slide_01, slide_executive_summary, slide_02,

    # 2. Context and business need
    slide_03,

    # 3. Architecture overview and platform constraints
    slide_04, slide_05, slide_06,

    # 4. End-to-end pipeline and produced artifacts
    slide_07, slide_08, slide_13,

    # 5. Preprocessing deep dive
    slide_10, slide_11, slide_12,
    slide_resolver_diagnostic_combined,

    # 6. ALFRED Vision integration
    slide_15, slide_16, slide_17,

    # 7. Delivery, quality gates, and upload
    slide_18, slide_stage_evaluation,

    # 8. ST ChatGPT / RAG configuration and routing
    slide_chatgpt_config_knobs, slide_system_prompt_and_routing,

    # 9. Value proposition, operations, and conclusion
    slide_artifact_added_value, slide_20, slide_19,
    slide_22,

    # Backup
    slide_23, slide_24,
]


def main():
    # Start from a clean PowerPoint package. Clearing all slides from an existing
    # deck can leave stale relationships in some templates, which PowerPoint may
    # report as an unreadable file.
    prs = Presentation()
    prs.slide_width = Cm(SLIDE_W)
    prs.slide_height = Cm(SLIDE_H)

    for builder in SLIDE_BUILDERS:
        builder(prs)

    output = save_presentation(prs)
    print(f"Saved: {output}")
    print(f"Slides: {len(prs.slides)}")
    print("Recommended timing: 25 main slides + 2 backup slides = about 40 minutes.")


if __name__ == "__main__":
    main()
