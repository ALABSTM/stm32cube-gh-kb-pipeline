from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "Atelier_Pratique_Pipeline_STM32Cube.pptx"
OUTPUT = ROOT / "Atelier_Pratique_Pipeline_STM32Cube_improved.pptx"

ST_BLUE = RGBColor(26, 35, 126)
ST_LIGHT_BLUE = RGBColor(218, 229, 255)
ST_ORANGE = RGBColor(255, 152, 0)
ST_GREEN = RGBColor(76, 175, 80)
ST_RED = RGBColor(210, 60, 60)
ST_GRAY = RGBColor(92, 96, 105)
LIGHT_GRAY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 20, 20)


def set_text(shape, text, size=14, color=BLACK, bold=False, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return tf


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Cm(0.8), Cm(0.35), Cm(24), Cm(1.1))
    set_text(box, title, size=23, color=ST_BLUE, bold=True, align=PP_ALIGN.LEFT)
    if subtitle:
        sub = slide.shapes.add_textbox(Cm(0.85), Cm(1.35), Cm(23.5), Cm(0.7))
        set_text(sub, subtitle, size=10.5, color=ST_GRAY, align=PP_ALIGN.LEFT)


def add_box(slide, x, y, w, h, text, fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=11, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    set_text(shape, text, size=font_size, color=BLACK, bold=bold)
    return shape


def add_label(slide, x, y, w, h, text, size=11, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    set_text(shape, text, size=size, color=color, bold=bold, align=align)
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=ST_BLUE):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(2)
    conn.line.end_arrowhead = True
    return conn


def add_bar(slide, x, y, label, value, color, max_value=100, width=13.5):
    add_label(slide, x, y - 0.05, 4.8, 0.45, label, size=10.5, bold=True)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + 5.1), Cm(y), Cm(width), Cm(0.45))
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(232, 236, 242)
    bg.line.color.rgb = RGBColor(232, 236, 242)
    fg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + 5.1), Cm(y), Cm(width * value / max_value), Cm(0.45))
    fg.fill.solid(); fg.fill.fore_color.rgb = color
    fg.line.color.rgb = color
    add_label(slide, x + 5.3 + width, y - 0.05, 2.2, 0.45, f"{value:.1f}%", size=10.5, bold=True, color=color)


def add_note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def move_last_slide_to(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    slides = list(sld_id_lst)
    last = slides[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(index, last)


def add_overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Start with the Big Picture", "Frame the workshop before entering commands: source → pipeline → KB → answer")

    add_box(slide, 0.8, 2.3, 3.4, 1.4, "GitHub\n25 series\n200+ repos", fill=RGBColor(232, 240, 255), font_size=11)
    add_box(slide, 5.0, 2.3, 3.4, 1.4, "Pipeline\n7 stages\n23 steps", fill=ST_LIGHT_BLUE, font_size=11)
    add_box(slide, 9.2, 2.3, 3.4, 1.4, "ST-ready JSON\n4 artifact types", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=11)
    add_box(slide, 13.4, 2.3, 3.4, 1.4, "KB #793\n38 datasources", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=11)
    add_box(slide, 17.6, 2.3, 3.4, 1.4, "ST ChatGPT\nGPT-5.1", fill=RGBColor(235, 232, 255), line=RGBColor(100, 85, 190), font_size=11)
    add_box(slide, 21.8, 2.3, 2.8, 1.4, "Sourced\nanswer", fill=RGBColor(230, 250, 230), line=ST_GREEN, font_size=11)

    for x in [4.25, 8.45, 12.65, 16.85, 21.05]:
        add_arrow(slide, x, 3.0, x + 0.65, 3.0)

    add_label(slide, 1.2, 5.1, 22.5, 0.8, "What the audience must understand first", size=16, color=ST_BLUE, bold=True)
    add_box(slide, 1.2, 6.1, 6.6, 2.1, "Modular\nEach stage can be replayed independently\n--skip-ingestion, --skip-chunking", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11, bold=False)
    add_box(slide, 9.2, 6.1, 6.6, 2.1, "Traceable\nEvery final card keeps source metadata\nrepo, issue, file path, URL", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11, bold=False)
    add_box(slide, 17.2, 6.1, 6.6, 2.1, "Production-ready\nSchema gates, upload IDs, retries,\nAlfred/NovaX safeguards", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11, bold=False)

    add_note(slide, "Open the atelier with this slide. The goal is not to start with commands but to put everyone in the same mental model: raw GitHub data becomes KB-ready evidence.")
    return slide


def add_image_problem_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Why Images Needed a Dedicated Solution", "Screenshots and PDF figures carry diagnostic signal, but text search cannot see pixels")

    add_box(slide, 1.0, 2.2, 10.5, 3.4, "Before Alfred\n\nIssue text says: \"see screenshot below\"\n[IMAGE ATTACHED]\n\nProblem: retrieval sees only a placeholder", fill=RGBColor(255, 235, 235), line=ST_RED, font_size=13)
    add_box(slide, 13.9, 2.2, 10.5, 3.4, "After Alfred\n\n[IMAGE DESCRIPTION] The screenshot shows STM32CubeIDE build warnings, GCC flags, tool settings, and version context.\n\nResult: visual content becomes searchable", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_arrow(slide, 11.7, 3.9, 13.6, 3.9, color=ST_ORANGE)
    add_label(slide, 11.6, 3.15, 2.3, 0.6, "Alfred\nVision", size=12, color=ST_ORANGE, bold=True, align=PP_ALIGN.CENTER)

    add_label(slide, 1.2, 7.1, 22.5, 0.65, "Measured confidence, not just API success", size=16, color=ST_BLUE, bold=True)
    add_bar(slide, 1.2, 8.2, "Issue images", 98.5, ST_GREEN)
    add_bar(slide, 1.2, 9.2, "PDF figures", 41.7, ST_ORANGE)
    add_bar(slide, 1.2, 10.2, "Total Alfred", 67.4, RGBColor(70, 120, 210))
    add_label(slide, 1.2, 11.4, 22.5, 1.2, "Interpretation: issue screenshots are strongly grounded in issue text; PDF figures are useful but harder to validate automatically because page context is often weak.", size=12, color=ST_GRAY)

    add_note(slide, "Use this before the existing ALFRED section. Emphasize that the project distinguishes operational success from contextual confidence.")
    return slide


def add_kb_novax_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "KB-First Agent Architecture: Alfred, KB and NovaX", "The fallback must never short-circuit the knowledge base")

    add_box(slide, 1.0, 2.2, 4.4, 1.2, "Engineer\nquestion", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_box(slide, 7.0, 2.0, 5.2, 1.6, "ST GitHub Analyzer\nPersona\nGPT-5.1", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_box(slide, 15.0, 1.1, 5.2, 1.55, "KB Tool #793\nALWAYS FIRST\nHybrid Search", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=12)
    add_box(slide, 15.0, 4.0, 5.2, 1.55, "NovaX Tool\nONLY IF KB EMPTY\nBugzilla + Community", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_box(slide, 7.0, 6.7, 5.2, 1.35, "Pipeline\n7 stages + upload", fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=12)

    add_arrow(slide, 5.5, 2.8, 6.9, 2.8)
    add_arrow(slide, 12.3, 2.7, 14.8, 1.9, color=ST_GREEN)
    add_arrow(slide, 12.3, 3.0, 14.8, 4.8, color=ST_ORANGE)
    add_arrow(slide, 9.6, 6.65, 17.5, 2.75, color=ST_BLUE)

    add_label(slide, 1.1, 9.0, 10.5, 2.0, "Why this matters:\n• If fallback is too permissive, the persona calls NovaX instead of the KB\n• This loses GitHub issue citations and reduces Quality Index\n• The restrictive tool description protects source-grounded answers", size=12, align=PP_ALIGN.LEFT)
    add_label(slide, 13.0, 9.0, 10.5, 2.0, "Current KB setup:\n• Search type: Hybrid\n• Semantic threshold: 0.55\n• Semantic docs: 25\n• Full-text docs: 12\n• Datasources: series × artifact type", size=12, align=PP_ALIGN.LEFT)
    return slide


def add_v3_cleaning_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Cleaning V3: Type-Aware Cleaning", "The key improvement between raw files and usable documents")

    add_box(slide, 1.0, 2.0, 4.0, 1.0, "raw_files_*.json", fill=LIGHT_GRAY, line=ST_GRAY, font_size=12)
    add_arrow(slide, 5.2, 2.5, 6.4, 2.5)
    add_box(slide, 6.5, 1.85, 4.2, 1.3, "detect_file_type()", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)

    types = [
        (1.0, 5.0, "Release Notes\nHTML", "strip_html()\n-33%", ST_BLUE),
        (5.8, 5.0, "README\nMarkdown", "clean_markdown()\n-15%", ST_BLUE),
        (10.6, 5.0, "Source code\n.c / .h", "strip_license()\n-10%", ST_BLUE),
        (15.4, 5.0, "CMSIS header\n>100 Ko", "filter_cmsis()\n-96%", ST_ORANGE),
        (20.2, 5.0, "Other", "passthrough\n0%", ST_GRAY),
    ]
    for x, y, title, action, color in types:
        add_arrow(slide, 8.6, 3.2, x + 1.8, 4.85, color=color)
        add_box(slide, x, y, 3.9, 1.3, title, fill=RGBColor(235, 240, 252), line=color, font_size=11)
        add_box(slide, x, y + 1.8, 3.9, 1.15, action, fill=WHITE, line=color, font_size=11)

    add_box(slide, 1.0, 10.0, 10.7, 1.35, "Concrete example: stm32h743xx.h\n1,458,301 chars → 52,140 chars (-96.4%)", fill=RGBColor(255, 245, 225), line=ST_ORANGE, font_size=12)
    add_box(slide, 13.0, 10.0, 10.7, 1.35, "Result: clean_files_<repo>_v3.json\n+31% estimated signal density per chunk", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_note(slide, "Insert after the existing Stage 2 Cleaning slide. This clarifies exactly why V3 exists and where it changes the pipeline.")
    return slide


def add_clean_to_docs_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Cleaning → Enrichment: Where the docs_* Files Come From", "This is the missing bridge between cleaned data and ST-ready delivery")

    add_box(slide, 1.0, 2.0, 4.5, 1.2, "clean_issues_*.json\nclean_files_*_v3.json", fill=LIGHT_GRAY, line=ST_GRAY, font_size=11)
    add_arrow(slide, 5.7, 2.6, 7.0, 2.6)
    add_box(slide, 7.1, 1.3, 5.2, 2.6, "issues_to_docs_v2.py\nfiles_to_docs_v2.py\n\nAdds structure and domain metadata", fill=ST_LIGHT_BLUE, line=ST_BLUE, font_size=12)
    add_arrow(slide, 12.5, 2.6, 14.0, 2.6)
    add_box(slide, 14.1, 2.0, 4.7, 1.2, "docs_issues_*_v2.json\ndocs_files_*_v3.json", fill=RGBColor(225, 245, 235), line=ST_GREEN, font_size=11)
    add_arrow(slide, 18.9, 2.6, 20.1, 2.6)
    add_box(slide, 20.2, 2.0, 3.8, 1.2, "Delivery\nst_ready_*.json", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=11)

    add_label(slide, 1.1, 5.0, 10.7, 0.5, "Issue enrichment metadata", size=14, color=ST_BLUE, bold=True)
    meta = [("layer", "HAL / BSP / CMSIS / MW / APP"), ("component", "SPI / I2C / UART / DMA / USB"), ("severity", "high / medium / low"), ("board", "NUCLEO-H743ZI / STM32H750B-DK"), ("issue_kind", "bug_report / question / feature")]
    y = 5.8
    for k, v in meta:
        add_box(slide, 1.2, y, 3.0, 0.65, k, fill=RGBColor(232, 240, 255), line=ST_BLUE, font_size=10)
        add_label(slide, 4.5, y + 0.08, 8.5, 0.45, v, size=10.5)
        y += 0.75

    add_label(slide, 14.0, 5.0, 9.5, 0.5, "Measured H7 coverage", size=14, color=ST_BLUE, bold=True)
    add_bar(slide, 14.0, 5.8, "Layer", 88.2, ST_GREEN, width=6.0)
    add_bar(slide, 14.0, 6.7, "Severity", 100.0, ST_GREEN, width=6.0)
    add_bar(slide, 14.0, 7.6, "Component", 72.6, ST_ORANGE, width=6.0)
    add_bar(slide, 14.0, 8.5, "Board", 45.3, ST_RED, width=6.0)
    add_label(slide, 14.0, 9.6, 9.5, 0.9, "Board is low because many issues simply do not mention a board; this is a source-data limitation, not only a heuristic failure.", size=10.5, color=ST_GRAY)
    return slide


def add_alfred_confidence_graph_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "ALFRED Evaluation: Real Results and Interpretation", "Contextual trust separates mature issue screenshots from harder PDF figures")
    add_bar(slide, 2.0, 3.0, "Issue images", 98.5, ST_GREEN, width=12.0)
    add_bar(slide, 2.0, 4.3, "PDF figures", 41.7, ST_ORANGE, width=12.0)
    add_bar(slide, 2.0, 5.6, "Total Alfred", 67.4, RGBColor(70, 120, 210), width=12.0)

    add_box(slide, 2.0, 7.6, 9.5, 2.0, "Issue images\n200 evaluated, 3 low coherence\nStrongly grounded in issue context", fill=RGBColor(230, 248, 235), line=ST_GREEN, font_size=12)
    add_box(slide, 13.2, 7.6, 9.5, 2.0, "PDF figures\n242 evaluated, 141 low coherence\nUseful, but page context is weak", fill=RGBColor(255, 241, 218), line=ST_ORANGE, font_size=12)
    add_label(slide, 2.0, 10.7, 21.0, 0.9, "Do not call this accuracy. It is contextual trust: percentage of descriptions above the coherence alert threshold S_coh ≥ 0.12.", size=12, color=ST_GRAY, bold=True)
    return slide


def add_slide_after(prs, slide_func, after_one_based):
    slide_func(prs)
    move_last_slide_to(prs, after_one_based)


def main():
    prs = Presentation(INPUT)

    # Insert from bottom to top so original slide numbers remain stable.
    add_slide_after(prs, add_alfred_confidence_graph_slide, 45)  # after ALFRED coverage
    add_slide_after(prs, add_kb_novax_slide, 41)                 # after timing insight
    add_slide_after(prs, add_image_problem_slide, 38)            # after ALFRED section divider
    add_slide_after(prs, add_clean_to_docs_slide, 34)            # after enrichment V2
    add_slide_after(prs, add_v3_cleaning_slide, 33)              # after cleaning normalization
    add_slide_after(prs, add_overview_slide, 9)                  # after architecture overview divider

    prs.save(OUTPUT)
    print(f"Saved improved deck: {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
