"""Convert test suite v2 CSV to XLSX."""
import csv
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    raise SystemExit(1)

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "datasets" / "07_delivery" / "st_ready" / "all_series_test_suite_50q_v2.csv"
XLSX_PATH = CSV_PATH.with_suffix(".xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Test Suite 50Q v2"

# Styles
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

with open(CSV_PATH, encoding="utf-8") as f:
    reader = csv.reader(f, delimiter=";")
    for row_idx, row in enumerate(reader, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

# Column widths
ws.column_dimensions["A"].width = 90   # Question
ws.column_dimensions["B"].width = 100  # Answer
ws.column_dimensions["C"].width = 20   # Tags

# Freeze header
ws.freeze_panes = "A2"

wb.save(XLSX_PATH)
print(f"Created: {XLSX_PATH}")
print(f"Rows: {ws.max_row - 1} questions")
