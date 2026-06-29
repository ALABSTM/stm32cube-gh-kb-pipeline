"""
Generate a quality test suite (Question, Answer, Tags) for all configured STM32Cube series.

The series list is inferred from shared/config/config.json + shared/config/config_*.json
(excluding config_all_series.json).

Usage:
    python pipeline_Automation/test_suites/generate_quality_test_suite.py
"""

from __future__ import annotations

import json
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    raise SystemExit(1)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
CONFIG_DIR = PROJECT_ROOT / "shared" / "config"
OUTPUT_DIR = PROJECT_ROOT / "datasets" / "07_delivery" / "st_ready"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
XLSX_PATH = OUTPUT_DIR / "quality_test_suite_all_series.xlsx"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _iter_series_configs() -> list[Path]:
    paths: list[Path] = []

    default_cfg = CONFIG_DIR / "config.json"
    if default_cfg.exists():
        paths.append(default_cfg)

    for p in sorted(CONFIG_DIR.glob("config_*.json")):
        if p.name == "config_all_series.json":
            continue
        paths.append(p)

    return paths


def _find_main_repo(cfg: dict, fallback: str) -> str:
    repos = cfg.get("repos", [])
    if isinstance(repos, list):
        for repo in repos:
            r = str(repo).strip()
            if r.lower().startswith("stm32cube"):
                return r
        if repos:
            return str(repos[0]).strip() or fallback
    return fallback


def _series_tag_from_cfg(cfg_path: Path, cfg: dict) -> str:
    if cfg_path.name == "config.json":
        main = _find_main_repo(cfg, "STM32CubeH7")
        if main.lower().startswith("stm32cube"):
            return main[len("STM32Cube") :].upper() or "H7"
        return "H7"

    stem = cfg_path.stem  # config_wl3
    tag = stem.replace("config_", "", 1).upper()
    return tag


def _build_questions_for_series(tag: str, main_repo: str) -> list[dict[str, str]]:
    # 4 robust questions per series keeps runtime reasonable while covering common failure modes.
    templates = [
        (
            "My {main_repo} project crashes with HardFault when enabling DMA in interrupt mode. "
            "Polling mode works. What should I verify first in clock, memory alignment, and interrupt configuration?"
        ),
        (
            "After updating to a newer {main_repo} package release, a previously working peripheral init now fails. "
            "What is the best migration/debug checklist to identify breaking changes quickly?"
        ),
        (
            "On {main_repo}, I2C sometimes gets stuck BUSY after a communication error and the bus never recovers. "
            "What is a safe software recovery sequence without full MCU reset?"
        ),
        (
            "I need lower power consumption on {main_repo} while preserving wake-up reliability. "
            "Which low-power configuration points should be audited first (clocks, GPIO state, peripherals, wake-up sources)?"
        ),
    ]

    return [
        {
            "question": t.format(tag=tag, main_repo=main_repo),
            "answer": "",
            "tags": tag,
        }
        for t in templates
    ]


def build_questions() -> list[dict[str, str]]:
    all_questions: list[dict[str, str]] = []

    for cfg_path in _iter_series_configs():
        cfg = _load_json(cfg_path)
        tag = _series_tag_from_cfg(cfg_path, cfg)
        main_repo = _find_main_repo(cfg, f"STM32Cube{tag}")
        all_questions.extend(_build_questions_for_series(tag=tag, main_repo=main_repo))

    return all_questions


QUESTIONS = build_questions()


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quality Test Suite"

    header_fill = PatternFill(start_color="1A2382", end_color="1A2382", fill_type="solid")
    header_font = Font(bold=True, size=11, color="FFFFFF")
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["Question", "Answer", "Tags"]
    col_widths = [95, 70, 14]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[chr(64 + col_idx)].width = width

    for row_idx, q in enumerate(QUESTIONS, start=2):
        row_data = [q["question"], q["answer"], q["tags"]]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align
            cell.border = thin_border
            cell.font = Font(size=10)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(QUESTIONS) + 1}"

    wb.save(XLSX_PATH)

    tags = sorted({q["tags"] for q in QUESTIONS})
    print(f"Generated: {XLSX_PATH}")
    print(f"Total questions: {len(QUESTIONS)}")
    print(f"Series covered ({len(tags)}): {', '.join(tags)}")


if __name__ == "__main__":
    main()
