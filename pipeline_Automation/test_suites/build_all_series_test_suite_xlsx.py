"""
Convert all_series_test_suite_50q.csv to .xlsx format.

Usage:
    python pipeline_Automation/build_all_series_test_suite_xlsx.py
"""

import csv
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    raise SystemExit(1)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = PROJECT_ROOT / "datasets" / "07_delivery" / "st_ready" / "all_series_test_suite_50q.csv"
XLSX_PATH = PROJECT_ROOT / "datasets" / "07_delivery" / "st_ready" / "all_series_test_suite_50q.xlsx"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Suite 50Q All Series"

    # Styles
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Column widths
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 15

    # Read CSV
    with open(CSV_PATH, "r", encoding="utf-8") as f:
        reader = csv.reader(f, delimiter=";")
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value.strip('"'))
                cell.alignment = wrap_align
                cell.border = thin_border

                if row_idx == 1:
                    cell.font = header_font_white
                    cell.fill = header_fill
                else:
                    cell.font = Font(size=10)

    # Freeze header
    ws.freeze_panes = "A2"

    wb.save(XLSX_PATH)
    print(f"Created: {XLSX_PATH}")
    print(f"Total rows: {ws.max_row - 1} questions")


if __name__ == "__main__":
    main()
