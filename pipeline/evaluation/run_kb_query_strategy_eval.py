import argparse
import csv
import json
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
import urllib3

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pipeline_Automation.upload.Get_Persona_KBs import (  # noqa: E402
    apiKey,
    clientAppName,
    generate_token,
    get_ST_GitHub_Analyzer_kb,
    proxies,
    url,
)

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


DEFAULT_STRATEGY_MATRIX = PROJECT_ROOT / "docs" / "evaluation" / "kb_strategy_matrix.csv"
DEFAULT_OUTPUT_CSV = PROJECT_ROOT / "docs" / "evaluation" / "kb_query_strategy_comparison.csv"


def infer_query_type(question: str, tags: str = "") -> str:
    text = (question or "").lower()
    tag_text = (tags or "").lower()

    if any(x in text for x in ["error", "fault", "crash", "stuck", "nack", "unknown device"]) or "debug" in tag_text:
        return "error_message"
    if any(x in text for x in ["hal_", "api", "function", "register", "callback"]):
        return "api_symbol"
    return "how_to"


def load_questions_csv(path: Path) -> list[dict[str, str]]:
    questions: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            qid = (row.get("query_id") or "").strip()
            qtext = (row.get("query_text") or "").strip()
            if qid and qtext:
                tags = (row.get("tags") or "").strip()
                if not tags:
                    tags = (row.get("notes") or "").strip()
                out = dict(row)
                out["tags"] = tags
                if not out.get("query_type"):
                    out["query_type"] = infer_query_type(qtext, tags=tags)
                questions.append(out)
    return questions


def load_questions_xlsx(path: Path, sheet_name: str = "") -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to read .xlsx question sets") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []

    col_map: dict[str, int] = {}
    for idx, name in enumerate(header):
        if name is None:
            continue
        key = str(name).strip().lower()
        col_map[key] = idx

    q_col = None
    for key in ("question", "query", "prompt"):
        if key in col_map:
            q_col = col_map[key]
            break
    if q_col is None:
        raise ValueError("Could not find a question column in xlsx (expected: Question)")

    tags_col = None
    for key in ("tags", "tag", "category"):
        if key in col_map:
            tags_col = col_map[key]
            break

    out: list[dict[str, str]] = []
    counter = 1
    for row in rows:
        if q_col >= len(row):
            continue
        qtext = str(row[q_col] or "").strip()
        if not qtext:
            continue

        tags = ""
        if tags_col is not None and tags_col < len(row):
            tags = str(row[tags_col] or "").strip()

        qid = f"Q{counter:03d}"
        out.append(
            {
                "query_id": qid,
                "query_text": qtext,
                "query_type": infer_query_type(qtext, tags=tags),
                "expected_repo_or_card": "",
                "tags": tags,
                "notes": tags,
            }
        )
        counter += 1

    return out


def load_questions(path: Path, sheet_name: str = "") -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        return load_questions_xlsx(path, sheet_name=sheet_name)
    return load_questions_csv(path)


def load_strategy_matrix(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Strategy matrix not found: {path}")

    strategies: dict[str, dict[str, str]] = {}
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sid = (row.get("strategy_id") or "").strip()
            if sid:
                strategies[sid] = row
    return strategies


@dataclass
class EvalRow:
    run_id: str
    run_date: str
    strategy_id: str
    search_type: str
    semantic_threshold: str
    semantic_max_doc_count: str
    fulltext_max_doc_count: str
    query_id: str
    query_text: str
    query_type: str
    expected_repo_or_card: str
    tags: str
    has_relevant_doc_top1: str
    has_relevant_doc_top3: str
    results_count: str
    top1_score: str
    latency_seconds: str
    notes: str


def call_kb_query(
    question: str,
    kb_id: int,
    semantic_threshold: float,
    semantic_max_docs_count: int,
    fulltext_max_docs_count: int,
    tags: list[str],
    remote_user: str,
    max_retries: int,
    request_timeout: int,
    use_proxy: bool,
) -> tuple[list[dict[str, Any]], dict[str, Any], float]:
    last_error = ""

    for attempt in range(1, max_retries + 1):
        # API deployments may expect either seconds or milliseconds timestamps.
        for ts_mode in ("seconds", "milliseconds"):
            timestamp = int(time.time()) if ts_mode == "seconds" else int(time.time() * 1000)
            nonce = random.randint(0, 999999)
            token = generate_token(clientAppName, "kb", apiKey, timestamp, nonce)

            headers = {
                "Content-Type": "application/json",
                "stchatgpt-auth-token": token,
                "stchatgpt-auth-nonce": str(nonce),
            }

            payload: dict[str, Any] = {
                "version": 1,
                "clientAppName": clientAppName,
                "timestamp": timestamp,
                "service": "kb",
                "type": "kb-query",
                "kb": kb_id,
                "query": question,
                "semanticThreshold": semantic_threshold,
                "semanticMaxDocsCount": semantic_max_docs_count,
                "fullTextMaxDocsCount": fulltext_max_docs_count,
            }
            if tags:
                payload["tags"] = tags
            if remote_user:
                payload["remoteUser"] = remote_user

            t0 = time.perf_counter()
            try:
                response = requests.post(
                    url,
                    json=payload,
                    headers=headers,
                    verify=False,
                    timeout=request_timeout,
                    proxies=proxies if use_proxy else None,
                )
                latency = time.perf_counter() - t0

                if response.status_code >= 400:
                    last_error = f"HTTP {response.status_code}: {response.text[:300]}"
                    if attempt < max_retries and response.status_code in {408, 425, 429, 500, 502, 503, 504}:
                        time.sleep(2 ** attempt)
                        break
                    return [], {"error": last_error}, latency

                data = response.json()
                if isinstance(data, dict) and data.get("errorCode") not in (None, 0):
                    last_error = f"errorCode={data.get('errorCode')}, message={data.get('message')}"
                    # Retry once with alternate timestamp format when window validation fails.
                    msg = str(data.get("message") or "").lower()
                    if data.get("errorCode") == 2000 and "invalid timestamp" in msg and ts_mode == "seconds":
                        continue
                    if attempt < max_retries:
                        time.sleep(2 ** attempt)
                        break
                    return [], {"error": last_error, "response": data}, latency

                results = data.get("results", []) if isinstance(data, dict) else []
                if not isinstance(results, list):
                    results = []
                return results, data if isinstance(data, dict) else {}, latency

            except requests.exceptions.Timeout:
                latency = time.perf_counter() - t0
                last_error = "Request Timeout"
                if attempt < max_retries:
                    time.sleep(2 ** attempt)
                    break
                return [], {"error": last_error}, latency
            except Exception as exc:
                latency = time.perf_counter() - t0
                last_error = str(exc)
                if attempt < max_retries:
                    time.sleep(2 ** attempt)
                    break
                return [], {"error": last_error}, latency

    return [], {"error": last_error or "unknown error"}, 0.0


def repo_hit_flags(expected_repo: str, results: list[dict[str, Any]]) -> tuple[str, str]:
    repo = (expected_repo or "").strip().lower()
    if not repo:
        return "", ""

    top1_hit = "0"
    top3_hit = "0"
    if results:
        first = json.dumps(results[0], ensure_ascii=False).lower()
        top1_hit = "1" if repo in first else "0"
        top3_blob = "\n".join(json.dumps(r, ensure_ascii=False).lower() for r in results[:3])
        top3_hit = "1" if repo in top3_blob else "0"
    return top1_hit, top3_hit


def write_run_artifacts(run_dir: Path, rows: list[dict[str, Any]], manifest: dict[str, Any]) -> None:
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    raw_path = run_dir / "responses_raw.jsonl"
    with raw_path.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    answers_path = run_dir / "responses_flat.csv"
    with answers_path.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "query_id",
            "query_text",
            "latency_seconds",
            "results_count",
            "top1_score",
            "top1_excerpt",
            "error",
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in rows:
            results = r.get("results", [])
            first = results[0] if isinstance(results, list) and results else {}
            first_score = ""
            if isinstance(first, dict) and first.get("score") is not None:
                first_score = str(first.get("score"))
            first_excerpt = ""
            if isinstance(first, dict):
                first_excerpt = str(first.get("pageContent") or "")[:240].replace("\n", " ")

            writer.writerow(
                {
                    "query_id": r.get("query_id", ""),
                    "query_text": r.get("query_text", ""),
                    "latency_seconds": f"{r.get('latency_seconds', 0.0):.3f}",
                    "results_count": len(results) if isinstance(results, list) else 0,
                    "top1_score": first_score,
                    "top1_excerpt": first_excerpt,
                    "error": r.get("error", ""),
                }
            )


def upsert_evaluation_rows(output_csv: Path, new_rows: list[EvalRow]) -> None:
    fieldnames = [
        "run_id",
        "run_date",
        "strategy_id",
        "search_type",
        "semantic_threshold",
        "semantic_max_doc_count",
        "fulltext_max_doc_count",
        "query_id",
        "query_text",
        "query_type",
        "expected_repo_or_card",
        "tags",
        "has_relevant_doc_top1",
        "has_relevant_doc_top3",
        "results_count",
        "top1_score",
        "latency_seconds",
        "notes",
    ]

    existing: list[dict[str, str]] = []
    if output_csv.exists():
        with output_csv.open("r", encoding="utf-8", newline="") as f:
            existing = list(csv.DictReader(f))

    by_key = {
        (r.get("run_id", ""), r.get("strategy_id", ""), r.get("query_id", "")): r
        for r in existing
    }

    for nr in new_rows:
        key = (nr.run_id, nr.strategy_id, nr.query_id)
        by_key[key] = {
            "run_id": nr.run_id,
            "run_date": nr.run_date,
            "strategy_id": nr.strategy_id,
            "search_type": nr.search_type,
            "semantic_threshold": nr.semantic_threshold,
            "semantic_max_doc_count": nr.semantic_max_doc_count,
            "fulltext_max_doc_count": nr.fulltext_max_doc_count,
            "query_id": nr.query_id,
            "query_text": nr.query_text,
            "query_type": nr.query_type,
            "expected_repo_or_card": nr.expected_repo_or_card,
            "tags": nr.tags,
            "has_relevant_doc_top1": nr.has_relevant_doc_top1,
            "has_relevant_doc_top3": nr.has_relevant_doc_top3,
            "results_count": nr.results_count,
            "top1_score": nr.top1_score,
            "latency_seconds": nr.latency_seconds,
            "notes": nr.notes,
        }

    merged = list(by_key.values())
    merged.sort(key=lambda r: (r.get("run_id", ""), r.get("strategy_id", ""), r.get("query_id", "")))

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(merged)


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate KB strategies directly via service=kb,type=kb-query.")
    parser.add_argument("--questions", default="docs/evaluation/all_series_test_suite_10q_diversified_chatwithalfred.xlsx")
    parser.add_argument("--xlsx-sheet", default="")
    parser.add_argument("--output-csv", default=str(DEFAULT_OUTPUT_CSV))
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--strategy-id", required=True)
    parser.add_argument("--strategy-matrix", default=str(DEFAULT_STRATEGY_MATRIX))
    parser.add_argument("--kb-id", type=int, default=0, help="Optional explicit KB id. If not set, resolves ST GitHub Analyzer KB.")
    parser.add_argument("--remote-user", default="abir.gharbi@st.com")
    parser.add_argument("--max-retries", type=int, default=2)
    parser.add_argument("--request-timeout", type=int, default=90)
    parser.add_argument("--delay", type=float, default=1.0)
    parser.add_argument("--use-proxy", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument(
        "--default-tags",
        default="",
        help="Fallback tags when a question row has no tags. Comma or semicolon separated.",
    )
    args = parser.parse_args()

    def parse_tags(value: str) -> list[str]:
        if not value:
            return []
        raw = value.replace(";", ",")
        return [t.strip() for t in raw.split(",") if t.strip()]

    def is_uuid(value: str) -> bool:
        return bool(
            re.match(
                r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[1-5][0-9a-fA-F]{3}-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}$",
                value,
            )
        )

    default_tags = parse_tags(args.default_tags)

    qpath = Path(args.questions)
    if not qpath.exists():
        raise FileNotFoundError(f"Questions file not found: {qpath}")

    questions = load_questions(qpath, sheet_name=args.xlsx_sheet)
    if args.limit > 0:
        questions = questions[: args.limit]
    if not questions:
        raise ValueError("No valid questions to run.")

    strategy_matrix = load_strategy_matrix(Path(args.strategy_matrix))
    strategy = strategy_matrix.get(args.strategy_id)
    if strategy is None:
        raise ValueError(f"Unknown strategy_id '{args.strategy_id}' in {args.strategy_matrix}")

    search_type = (strategy.get("search_type") or "").strip()
    semantic_threshold = float((strategy.get("semantic_threshold") or "0.55").strip())
    semantic_max_doc_count = int((strategy.get("semantic_max_doc_count") or "20").strip())
    fulltext_max_doc_count = int((strategy.get("fulltext_max_doc_count") or "12").strip())

    kb_id = args.kb_id
    kb_meta: dict[str, Any] = {}
    if kb_id <= 0:
        kb = get_ST_GitHub_Analyzer_kb()
        if not kb:
            raise RuntimeError("Could not resolve ST GitHub Analyzer KB. Pass --kb-id explicitly.")
        kb_id = int(kb.get("id"))
        kb_meta = kb
    else:
        kb = get_ST_GitHub_Analyzer_kb()
        if kb and int(kb.get("id") or 0) == kb_id:
            kb_meta = kb

    tag_items: list[dict[str, Any]] = []
    kb_tags_def = kb_meta.get("kbTagsDef") if isinstance(kb_meta, dict) else None
    if isinstance(kb_tags_def, dict):
        items = kb_tags_def.get("items")
        if isinstance(items, list):
            tag_items = [i for i in items if isinstance(i, dict)]
    elif isinstance(kb_tags_def, list):
        tag_items = [i for i in kb_tags_def if isinstance(i, dict)]

    label_to_id: dict[str, str] = {}
    for item in tag_items:
        label = str(item.get("label") or "").strip().lower()
        tid = str(item.get("id") or "").strip()
        if label and tid:
            label_to_id[label] = tid

    print("\n=== KB query strategy under test ===")
    print(f"strategy_id: {args.strategy_id}")
    print(f"kb_id: {kb_id}")
    print(f"search_type: {search_type}")
    print(f"semantic_threshold: {semantic_threshold}")
    print(f"semantic_max_doc_count: {semantic_max_doc_count}")
    print(f"fulltext_max_doc_count: {fulltext_max_doc_count}")

    run_date = datetime.now().strftime("%Y-%m-%d")
    run_dir = PROJECT_ROOT / "docs" / "evaluation" / "runs_kb_query" / f"{args.run_id}_{args.strategy_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

    raw_rows: list[dict[str, Any]] = []
    eval_rows: list[EvalRow] = []

    for idx, q in enumerate(questions, start=1):
        qid = (q.get("query_id") or "").strip()
        qtext = (q.get("query_text") or "").strip()
        qtype = (q.get("query_type") or "").strip()
        expected_repo = (q.get("expected_repo_or_card") or "").strip()
        q_tags = parse_tags((q.get("tags") or "").strip())
        raw_tags = q_tags or default_tags
        effective_tags: list[str] = []
        unresolved: list[str] = []
        for raw_tag in raw_tags:
            t = raw_tag.strip()
            if not t:
                continue
            if is_uuid(t):
                effective_tags.append(t)
                continue
            mapped = label_to_id.get(t.lower())
            if mapped:
                effective_tags.append(mapped)
            else:
                unresolved.append(t)

        print(f"[{idx}/{len(questions)}] {qid}: {qtext}")
        if effective_tags:
            print(f"       tags={effective_tags}")
        if unresolved:
            print(f"       [WARN] unresolved tag labels ignored: {unresolved}")
        results, response_json, latency = call_kb_query(
            question=qtext,
            kb_id=kb_id,
            semantic_threshold=semantic_threshold,
            semantic_max_docs_count=semantic_max_doc_count,
            fulltext_max_docs_count=fulltext_max_doc_count,
            tags=effective_tags,
            remote_user=args.remote_user,
            max_retries=args.max_retries,
            request_timeout=args.request_timeout,
            use_proxy=args.use_proxy,
        )

        top1, top3 = repo_hit_flags(expected_repo, results)
        top_score = ""
        if results and isinstance(results[0], dict) and results[0].get("score") is not None:
            top_score = str(results[0].get("score"))

        note = ""
        err = ""
        if isinstance(response_json, dict):
            err = str(response_json.get("error") or "").strip()
        if err:
            note = f"API error: {err}"
        elif not results:
            note = "No results returned"

        raw_rows.append(
            {
                "query_id": qid,
                "query_text": qtext,
                "query_type": qtype,
                "expected_repo_or_card": expected_repo,
                "tags": ",".join(effective_tags),
                "latency_seconds": latency,
                "results": results,
                "response_json": response_json,
                "error": err,
            }
        )

        eval_rows.append(
            EvalRow(
                run_id=args.run_id,
                run_date=run_date,
                strategy_id=args.strategy_id,
                search_type=search_type,
                semantic_threshold=str(semantic_threshold),
                semantic_max_doc_count=str(semantic_max_doc_count),
                fulltext_max_doc_count=str(fulltext_max_doc_count),
                query_id=qid,
                query_text=qtext,
                query_type=qtype,
                expected_repo_or_card=expected_repo,
                tags=",".join(effective_tags),
                has_relevant_doc_top1=top1,
                has_relevant_doc_top3=top3,
                results_count=str(len(results)),
                top1_score=top_score,
                latency_seconds=f"{latency:.3f}",
                notes=note,
            )
        )

        if args.delay > 0:
            time.sleep(args.delay)

    manifest = {
        "run_id": args.run_id,
        "run_date": run_date,
        "strategy_id": args.strategy_id,
        "kb_id": kb_id,
        "strategy_params": {
            "search_type": search_type,
            "semantic_threshold": semantic_threshold,
            "semantic_max_doc_count": semantic_max_doc_count,
            "fulltext_max_doc_count": fulltext_max_doc_count,
        },
        "questions_file": str(qpath),
        "xlsx_sheet": args.xlsx_sheet,
        "remote_user": args.remote_user,
        "request_timeout": args.request_timeout,
        "max_retries": args.max_retries,
        "delay": args.delay,
        "use_proxy": args.use_proxy,
    }

    write_run_artifacts(run_dir, raw_rows, manifest)
    upsert_evaluation_rows(Path(args.output_csv), eval_rows)

    print("\nRun completed.")
    print(f"Artifacts: {run_dir}")
    print(f"Evaluation CSV updated: {args.output_csv}")


if __name__ == "__main__":
    main()
