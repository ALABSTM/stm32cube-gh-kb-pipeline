import argparse
import csv
import json
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
import urllib3

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pipeline_Automation.upload.Get_Persona_KBs import (
    apiKey,
    clientAppName,
    generate_token,
    proxies,
    url,
)

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


CHAT_SERVICE = "chat"
DEFAULT_PERSONA = "ST GitHub Analyzer"
DEFAULT_STRATEGY_MATRIX = PROJECT_ROOT / "docs" / "evaluation" / "kb_strategy_matrix.csv"


def is_proxy_tunnel_error(message: str) -> bool:
    text = (message or "").lower()
    return (
        "proxyerror" in text
        or "unable to connect to proxy" in text
        or "tunnel connection failed" in text
        or "502 bad gateway" in text
    )


def infer_query_type(question: str, tags: str = "") -> str:
    text = (question or "").lower()
    tag_text = (tags or "").lower()

    if any(x in text for x in ["error", "fault", "crash", "stuck", "nack", "unknown device"]) or "debug" in tag_text:
        return "error_message"
    if any(x in text for x in ["hal_", "api", "function", "register", "callback"]):
        return "api_symbol"
    return "how_to"


def load_questions_csv(path: Path) -> list[dict[str, str]]:
    questions: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            qid = (row.get("query_id") or "").strip()
            qtext = (row.get("query_text") or "").strip()
            if qid and qtext:
                questions.append(row)
    return questions


def load_questions_xlsx(path: Path, sheet_name: str = "") -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to read .xlsx question sets") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []

    col_map: dict[str, int] = {}
    for idx, name in enumerate(header):
        if name is None:
            continue
        key = str(name).strip().lower()
        col_map[key] = idx

    # Common expected headers in the test suite files.
    q_col = None
    for key in ("question", "query", "prompt"):
        if key in col_map:
            q_col = col_map[key]
            break
    if q_col is None:
        raise ValueError("Could not find a question column in xlsx (expected: Question)")

    tags_col = None
    for key in ("tags", "tag", "category"):
        if key in col_map:
            tags_col = col_map[key]
            break

    out: list[dict[str, str]] = []
    counter = 1
    for row in rows:
        if q_col >= len(row):
            continue
        qtext = str(row[q_col] or "").strip()
        if not qtext:
            continue

        tags = ""
        if tags_col is not None and tags_col < len(row):
            tags = str(row[tags_col] or "").strip()

        qid = f"Q{counter:03d}"
        out.append(
            {
                "query_id": qid,
                "query_text": qtext,
                "query_type": infer_query_type(qtext, tags=tags),
                "expected_repo_or_card": "",
                "expected_keywords": "",
                "keywords_mode": "any",
                "forbidden_regex": "",
                "notes": tags,
            }
        )
        counter += 1

    return out


def load_questions(path: Path, sheet_name: str = "") -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        return load_questions_xlsx(path, sheet_name=sheet_name)
    return load_questions_csv(path)


@dataclass
class EvalRow:
    run_id: str
    run_date: str
    strategy_id: str
    persona: str
    search_type: str
    semantic_threshold: str
    semantic_max_doc_count: str
    fulltext_max_doc_count: str
    query_id: str
    query_text: str
    query_type: str
    expected_repo_or_card: str
    has_relevant_doc_top1: str
    has_relevant_doc_top3: str
    final_answer_correct: str
    hallucination: str
    latency_seconds: str
    notes: str


def load_strategy_matrix(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Strategy matrix not found: {path}")

    strategies: dict[str, dict[str, str]] = {}
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sid = (row.get("strategy_id") or "").strip()
            if sid:
                strategies[sid] = row
    return strategies


def extract_text(response_json: dict[str, Any]) -> str:
    if not isinstance(response_json, dict):
        return ""

    completion = response_json.get("completion")
    if isinstance(completion, str) and completion.strip():
        return completion.strip()

    for key in ("text", "content", "message", "answer", "final"):
        val = response_json.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()

    choices = response_json.get("choices")
    if isinstance(choices, list) and choices:
        first = choices[0]
        if isinstance(first, dict):
            msg = first.get("message")
            if isinstance(msg, dict):
                content = msg.get("content")
                if isinstance(content, str) and content.strip():
                    return content.strip()
            for key in ("text", "content"):
                val = first.get(key)
                if isinstance(val, str) and val.strip():
                    return val.strip()

    return ""


def extract_sources(response_json: Any) -> list[str]:
    sources: list[str] = []

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                lk = str(key).lower()
                if lk in {"sources", "documents", "citations", "evidence", "context"} and isinstance(value, list):
                    for item in value:
                        if isinstance(item, dict):
                            parts: list[str] = []
                            for k in ("title", "name", "url", "externalURL", "repo", "path", "id", "label"):
                                v = item.get(k)
                                if v is not None:
                                    parts.append(str(v))
                            text = " | ".join(parts).strip()
                            if text:
                                sources.append(text)
                        elif isinstance(item, str) and item.strip():
                            sources.append(item.strip())
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(response_json)

    # preserve order, remove duplicates
    dedup: list[str] = []
    seen = set()
    for s in sources:
        if s not in seen:
            dedup.append(s)
            seen.add(s)
    return dedup


def call_persona(
    question: str,
    persona: str,
    remote_user: str,
    temperature: float,
    use_proxy: bool,
    max_retries: int,
    request_timeout: int,
    max_response_tokens: int,
) -> tuple[str, dict[str, Any], float]:
    last_error = ""
    effective_use_proxy = use_proxy
    proxy_fallback_used = False
    for attempt in range(1, max_retries + 1):
        timestamp = int(time.time())
        nonce = random.randint(0, 999999)
        token = generate_token(clientAppName, CHAT_SERVICE, apiKey, timestamp, nonce)

        headers = {
            "Content-Type": "application/json",
            "stchatgpt-auth-token": token,
            "stchatgpt-auth-nonce": str(nonce),
        }
        payload = {
            "version": 1,
            "clientAppName": clientAppName,
            "service": CHAT_SERVICE,
            "timestamp": timestamp,
            "messages": [{"role": "user", "content": question}],
            "temperature": temperature,
            "persona": persona,
            "maxResponseTokens": max_response_tokens,
            "responseFormat": "text",
        }
        if remote_user:
            payload["remoteUser"] = remote_user

        t0 = time.perf_counter()
        try:
            response = requests.post(
                url,
                json=payload,
                headers=headers,
                verify=False,
                timeout=request_timeout,
                proxies=proxies if effective_use_proxy else None,
            )
            latency = time.perf_counter() - t0
            if response.status_code >= 400:
                last_error = f"HTTP {response.status_code}: {response.text[:300]}"
                if effective_use_proxy and response.status_code == 502:
                    # Retry once without proxy when tunnel/proxy returns bad gateway.
                    effective_use_proxy = False
                    proxy_fallback_used = True
                    if attempt < max_retries:
                        continue
                if attempt < max_retries and response.status_code in {408, 425, 429, 500, 502, 503, 504}:
                    time.sleep(2 ** attempt)
                    continue
                return "", {"error": last_error}, latency

            data = response.json()
            text = extract_text(data)
            if proxy_fallback_used and isinstance(data, dict):
                data.setdefault("warning", "proxy_fallback_no_proxy")
            return text, data, latency

        except requests.exceptions.Timeout:
            latency = time.perf_counter() - t0
            last_error = "Request Timeout"
            if attempt < max_retries:
                time.sleep(2 ** attempt)
                continue
            return "Request Timeout", {"error": last_error}, latency
        except Exception as exc:
            latency = time.perf_counter() - t0
            last_error = str(exc)
            if effective_use_proxy and is_proxy_tunnel_error(last_error):
                # Retry transparently without proxy for unstable proxy tunnel errors.
                effective_use_proxy = False
                proxy_fallback_used = True
                if attempt < max_retries:
                    continue
            if attempt < max_retries:
                time.sleep(2 ** attempt)
                continue
            return "", {"error": last_error}, latency

    return "", {"error": last_error or "unknown error"}, 0.0


def match_expected(answer_text: str, expected_keywords: str, mode: str) -> str:
    expected_keywords = (expected_keywords or "").strip()
    if not expected_keywords:
        return ""
    parts = [p.strip() for p in expected_keywords.split("|") if p.strip()]
    if not parts:
        return ""

    flags = re.IGNORECASE
    if mode.lower() == "all":
        ok = all(re.search(p, answer_text, flags=flags) for p in parts)
    else:
        ok = any(re.search(p, answer_text, flags=flags) for p in parts)
    return "1" if ok else "0"


def detect_hallucination(answer_text: str, forbidden_regex: str) -> str:
    forbidden_regex = (forbidden_regex or "").strip()
    if not forbidden_regex:
        return ""
    try:
        return "1" if re.search(forbidden_regex, answer_text, flags=re.IGNORECASE) else "0"
    except re.error:
        return ""


def write_run_artifacts(run_dir: Path, rows: list[dict[str, Any]], manifest: dict[str, Any]) -> None:
    run_dir.mkdir(parents=True, exist_ok=True)

    manifest_path = run_dir / "run_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    raw_path = run_dir / "responses_raw.jsonl"
    with raw_path.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    answers_path = run_dir / "responses_flat.csv"
    with answers_path.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "query_id", "query_text", "latency_seconds", "answer_text", "sources_top1", "sources_top3", "error"
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in rows:
            src = r.get("sources", [])
            writer.writerow(
                {
                    "query_id": r.get("query_id", ""),
                    "query_text": r.get("query_text", ""),
                    "latency_seconds": f"{r.get('latency_seconds', 0.0):.3f}",
                    "answer_text": (r.get("answer_text", "") or "").replace("\n", " "),
                    "sources_top1": src[0] if len(src) >= 1 else "",
                    "sources_top3": " || ".join(src[:3]) if src else "",
                    "error": r.get("error", ""),
                }
            )


def upsert_evaluation_rows(output_csv: Path, new_rows: list[EvalRow]) -> None:
    fieldnames = [
        "run_id",
        "run_date",
        "strategy_id",
        "persona",
        "search_type",
        "semantic_threshold",
        "semantic_max_doc_count",
        "fulltext_max_doc_count",
        "query_id",
        "query_text",
        "query_type",
        "expected_repo_or_card",
        "has_relevant_doc_top1",
        "has_relevant_doc_top3",
        "final_answer_correct",
        "hallucination",
        "latency_seconds",
        "notes",
    ]

    existing: list[dict[str, str]] = []
    if output_csv.exists():
        with output_csv.open("r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                existing.append(row)

    by_key = {
        (r.get("run_id", ""), r.get("strategy_id", ""), r.get("query_id", "")): r
        for r in existing
    }

    for nr in new_rows:
        key = (nr.run_id, nr.strategy_id, nr.query_id)
        by_key[key] = {
            "run_id": nr.run_id,
            "run_date": nr.run_date,
            "strategy_id": nr.strategy_id,
            "persona": nr.persona,
            "search_type": nr.search_type,
            "semantic_threshold": nr.semantic_threshold,
            "semantic_max_doc_count": nr.semantic_max_doc_count,
            "fulltext_max_doc_count": nr.fulltext_max_doc_count,
            "query_id": nr.query_id,
            "query_text": nr.query_text,
            "query_type": nr.query_type,
            "expected_repo_or_card": nr.expected_repo_or_card,
            "has_relevant_doc_top1": nr.has_relevant_doc_top1,
            "has_relevant_doc_top3": nr.has_relevant_doc_top3,
            "final_answer_correct": nr.final_answer_correct,
            "hallucination": nr.hallucination,
            "latency_seconds": nr.latency_seconds,
            "notes": nr.notes,
        }

    merged = list(by_key.values())
    merged.sort(key=lambda r: (r.get("run_id", ""), r.get("strategy_id", ""), r.get("query_id", "")))

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(merged)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run persona API on a question set and update evaluation CSV.")
    parser.add_argument("--questions", default="docs/evaluation/all_series_test_suite_50q.xlsx")
    parser.add_argument("--xlsx-sheet", default="", help="Optional sheet name when --questions is an xlsx file")
    parser.add_argument("--output-csv", default="docs/evaluation/kb_strategy_comparison_template.csv")
    parser.add_argument("--run-id", required=True, help="Run identifier, e.g. R0")
    parser.add_argument("--strategy-id", required=True, help="Strategy identifier, e.g. S0_CURRENT")
    parser.add_argument("--persona", default=DEFAULT_PERSONA)
    parser.add_argument("--strategy-matrix", default=str(DEFAULT_STRATEGY_MATRIX))
    parser.add_argument("--remote-user", default="abir.gharbi@st.com")
    parser.add_argument("--temperature", type=float, default=0.3)
    parser.add_argument("--max-retries", type=int, default=3)
    parser.add_argument("--request-timeout", type=int, default=90)
    parser.add_argument("--max-response-tokens", type=int, default=512)
    parser.add_argument("--delay", type=float, default=1.0, help="Delay in seconds between requests")
    parser.add_argument("--use-proxy", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()

    qpath = Path(args.questions)
    if not qpath.exists():
        raise FileNotFoundError(f"Questions file not found: {qpath}")

    questions: list[dict[str, str]] = load_questions(qpath, sheet_name=args.xlsx_sheet)

    if args.limit > 0:
        questions = questions[: args.limit]

    if not questions:
        raise ValueError("No valid questions to run.")

    strategy_matrix = load_strategy_matrix(Path(args.strategy_matrix))
    strategy = strategy_matrix.get(args.strategy_id)
    if strategy is None:
        raise ValueError(
            f"Unknown strategy_id '{args.strategy_id}' in {args.strategy_matrix}. "
            "Define the strategy first in kb_strategy_matrix.csv."
        )

    search_type = (strategy.get("search_type") or "").strip()
    semantic_threshold = (strategy.get("semantic_threshold") or "").strip()
    semantic_max_doc_count = (strategy.get("semantic_max_doc_count") or "").strip()
    fulltext_max_doc_count = (strategy.get("fulltext_max_doc_count") or "").strip()

    print("\n=== Strategy under test ===")
    print(f"strategy_id: {args.strategy_id}")
    print(f"persona: {args.persona}")
    print(f"search_type: {search_type}")
    print(f"semantic_threshold: {semantic_threshold}")
    print(f"semantic_max_doc_count: {semantic_max_doc_count}")
    print(f"fulltext_max_doc_count: {fulltext_max_doc_count}")
    print(
        "WARNING: this script does not change KB backend parameters automatically. "
        "Before running, ensure the ST GitHub Analyzer persona is configured with exactly these values."
    )

    run_date = datetime.now().strftime("%Y-%m-%d")
    run_dir = Path("docs/evaluation/runs") / f"{args.run_id}_{args.strategy_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

    raw_rows: list[dict[str, Any]] = []
    eval_rows: list[EvalRow] = []

    for idx, q in enumerate(questions, start=1):
        qid = (q.get("query_id") or "").strip()
        qtext = (q.get("query_text") or "").strip()
        qtype = (q.get("query_type") or "").strip()
        expected_repo = (q.get("expected_repo_or_card") or "").strip()
        expected_keywords = (q.get("expected_keywords") or "").strip()
        keywords_mode = (q.get("keywords_mode") or "any").strip().lower() or "any"
        forbidden_regex = (q.get("forbidden_regex") or "").strip()

        print(f"[{idx}/{len(questions)}] {qid}: {qtext}")
        answer_text, response_json, latency = call_persona(
            question=qtext,
            persona=args.persona,
            remote_user=args.remote_user,
            temperature=args.temperature,
            use_proxy=args.use_proxy,
            max_retries=args.max_retries,
            request_timeout=args.request_timeout,
            max_response_tokens=args.max_response_tokens,
        )

        sources = extract_sources(response_json)

        top1 = "0"
        top3 = "0"
        if expected_repo:
            repo_norm = expected_repo.lower()
            top1 = "1" if len(sources) >= 1 and repo_norm in sources[0].lower() else "0"
            top3 = "1" if any(repo_norm in s.lower() for s in sources[:3]) else "0"

        final_correct = match_expected(answer_text, expected_keywords, keywords_mode) or "0"
        hallucination = detect_hallucination(answer_text, forbidden_regex) or "0"

        if not hallucination and answer_text.strip().lower() == "request timeout":
            hallucination = "0"

        notes = ""
        if not answer_text:
            notes = "No answer extracted from API response"
        elif not sources:
            notes = "No explicit source/citation list returned by API"

        if isinstance(response_json, dict) and response_json.get("error"):
            notes = f"{notes} | API error: {response_json.get('error')}".strip(" |")

        raw_rows.append(
            {
                "query_id": qid,
                "query_text": qtext,
                "query_type": qtype,
                "expected_repo_or_card": expected_repo,
                "answer_text": answer_text,
                "latency_seconds": latency,
                "sources": sources,
                "response_json": response_json,
                "error": response_json.get("error", "") if isinstance(response_json, dict) else "",
            }
        )

        if args.delay > 0:
            time.sleep(args.delay)

        eval_rows.append(
            EvalRow(
                run_id=args.run_id,
                run_date=run_date,
                strategy_id=args.strategy_id,
                persona=args.persona,
                search_type=search_type,
                semantic_threshold=semantic_threshold,
                semantic_max_doc_count=semantic_max_doc_count,
                fulltext_max_doc_count=fulltext_max_doc_count,
                query_id=qid,
                query_text=qtext,
                query_type=qtype,
                expected_repo_or_card=expected_repo,
                has_relevant_doc_top1=top1,
                has_relevant_doc_top3=top3,
                final_answer_correct=final_correct,
                hallucination=hallucination,
                latency_seconds=f"{latency:.3f}",
                notes=notes,
            )
        )

    manifest = {
        "run_id": args.run_id,
        "run_date": run_date,
        "strategy_id": args.strategy_id,
        "persona": args.persona,
        "strategy_params": {
            "search_type": search_type,
            "semantic_threshold": semantic_threshold,
            "semantic_max_doc_count": semantic_max_doc_count,
            "fulltext_max_doc_count": fulltext_max_doc_count,
        },
        "questions_file": str(qpath),
        "xlsx_sheet": args.xlsx_sheet,
        "remote_user": args.remote_user,
        "temperature": args.temperature,
        "max_retries": args.max_retries,
        "request_timeout": args.request_timeout,
        "max_response_tokens": args.max_response_tokens,
        "delay": args.delay,
        "use_proxy": args.use_proxy,
        "note": "Backend KB parameters must be set manually to match strategy_params before this run.",
    }

    write_run_artifacts(run_dir, raw_rows, manifest)
    upsert_evaluation_rows(Path(args.output_csv), eval_rows)

    print("\nRun completed.")
    print(f"Artifacts: {run_dir}")
    print(f"Evaluation CSV updated: {args.output_csv}")


if __name__ == "__main__":
    main()
